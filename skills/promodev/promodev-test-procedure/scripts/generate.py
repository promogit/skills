# -*- coding: utf-8 -*-
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU

import os
HERE=os.path.dirname(os.path.abspath(__file__))
IMG=os.path.join(HERE,"..","assets","images")+os.sep
NAVY="1B2A4A"; ACCENT="E8683C"; BLUE="3B5BA5"; ZEBRA="F5F7FB"; TINT="EAF0FB"
GREY="EDEFF3"; YELLOW="FFF6E6"; YELLOWB="F3C892"; GREENF="D4EDBC"; GREENT="38761D"
REDF="F4CCCC"; REDT="A61C00"; NEUTF="E6E6E6"; WHITE="FFFFFF"; LINE="D5DAE3"
FONT="Calibri"
INFO={}   # valeurs réelles de l'onglet 1, par libellé exact (rempli en main)
def F(sz=11,b=False,color="26324A",it=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=it)
def fill(c): return PatternFill("solid",fgColor=c)
thin=Side(style="thin",color=LINE); border=Border(left=thin,right=thin,top=thin,bottom=thin)
def bd(ws,coord): ws[coord].border=border
wrap=Alignment(wrap_text=True,vertical="top")
wrapv=Alignment(wrap_text=True,vertical="center")
wrapc=Alignment(wrap_text=True,vertical="center",horizontal="center")
center=Alignment(vertical="center",horizontal="center")
leftc=Alignment(vertical="center",horizontal="left",indent=1)

def title_bar(ws,c1,c2,row,title,sub):
    ws.merge_cells(start_row=row,start_column=c1,end_row=row,end_column=c2)
    c=ws.cell(row,c1,title); c.font=F(18,True,WHITE); c.fill=fill(NAVY); c.alignment=Alignment(vertical="center",indent=1)
    ws.row_dimensions[row].height=42
    ws.merge_cells(start_row=row+1,start_column=c1,end_row=row+1,end_column=c2)
    for cc in range(c1,c2+1): ws.cell(row+1,cc).fill=fill(ACCENT)
    ws.row_dimensions[row+1].height=5
    ws.merge_cells(start_row=row+2,start_column=c1,end_row=row+2,end_column=c2)
    s=ws.cell(row+2,c1,sub); s.font=F(11,False,"6B7280"); s.alignment=Alignment(indent=1,vertical="center")
    ws.row_dimensions[row+2].height=22
    return row+3

def section(ws,col,row,text):
    c=ws.cell(row,col,"  "+text); c.font=F(13,True,NAVY); c.fill=fill(TINT); c.alignment=leftc
    ws.cell(row,col-1).fill=fill(ACCENT)
    return row

def place_center(ws,path,col0,row_idx,colw_chars,target_w,max_h=210):
    """Place une image centrée dans la cellule ; renvoie la hauteur de ligne (pts) à appliquer."""
    im=XLImage(path); ratio=im.height/im.width
    w=target_w; h=int(w*ratio)
    if h>max_h: h=max_h; w=int(h/ratio)
    im.width=w; im.height=h
    row_h_pts=h*0.75+22
    cell_px_w=int(colw_chars*7)+5
    cell_px_h=int(row_h_pts*4/3)
    offx=max(0,(cell_px_w-w)//2); offy=max(0,(cell_px_h-h)//2)
    marker=AnchorMarker(col=col0,colOff=pixels_to_EMU(offx),row=row_idx-1,rowOff=pixels_to_EMU(offy))
    im.anchor=OneCellAnchor(_from=marker,ext=XDRPositiveSize2D(pixels_to_EMU(w),pixels_to_EMU(h)))
    ws.add_image(im)
    return row_h_pts

from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D as _XSZ
from openpyxl.utils import get_column_letter as _gcl

def _col_px(g,ci):
    w=g.column_dimensions[_gcl(ci)].width or 8.43
    return int(w*7)+5
def _y_emu_above(g,r):
    tot=0.0
    for i in range(1,r):
        h=g.row_dimensions[i].height
        tot+=(h if h is not None else 15.0)
    return int(tot*12700)
def place_abs(g,filename,col1,row1,target_w,max_h,span,minh):
    """Ancre l'image en position ABSOLUE (taille + position figées, identiques partout)."""
    im=XLImage(IMG+filename); ratio=im.height/im.width
    w=target_w; h=int(w*ratio)
    if h>max_h: h=max_h; w=int(h/ratio)
    rh_pts=max(h*0.75+22,minh); g.row_dimensions[row1].height=rh_pts
    x_left=sum(_col_px(g,c) for c in range(1,col1))
    zone_w=sum(_col_px(g,c) for c in range(col1,col1+span))
    offx=max(0,(zone_w-w)//2)
    row_px=int(rh_pts*4/3); offy=max(0,(row_px-h)//2)
    X=int((x_left+offx)*9525); Y=_y_emu_above(g,row1)+int(offy*9525)
    im.width=w; im.height=h
    im.anchor=AbsoluteAnchor(pos=XDRPoint2D(X,Y),ext=_XSZ(int(w*9525),int(h*9525)))
    g.add_image(im)
    return rh_pts

GAME_META={
 "IG":  {"name":"INSTANT GAGNANT","file":"IG","type_val":"Instant gagnant","tirage":False,
         "mech":"résultat immédiat gagné / perdu"},
 "TAS": {"name":"TIRAGE AU SORT","file":"TAS","type_val":"Tirage au sort","tirage":True,
         "mech":"désignation des gagnants au tirage, pas de résultat immédiat"},
 "AUTO":{"name":"100 % GAGNANT","file":"100pct gagnant","type_val":"100 % gagnant (instant gagnant + lot de consolation)","tirage":False,
         "mech":"100 % gagnant : gros lot pour les instants gagnants, lot de consolation pour les autres"},
}
OBL_META={
 "AOA":{"name":"avec obligation d'achat","file":"avec achat","just":True},
 "SOA":{"name":"sans obligation d'achat (SOA)","file":"sans achat SOA","just":False},
}
CATCOL={"Accès & états":"E7EEF9","Home page":"E9F5EC","Formulaire":"FBF0E4","Preuve d'achat":"F3EAF7",
 "Participation SOA":"F3EAF7","Jeu · Instant gagnant":"FDECEC","Jeu · Gros lot (IG)":"FDECEC",
 "Jeu · Lot de consolation":"FCE9D6","Jeu · 100 % gagnant":"FDECEC","Jeu · Tirage au sort":"E4F3F5",
 "Dotation / gain":"FFF7E0","Consentements":"F0ECF9","Anti-fraude":"ECEFF2","E-mails":"E9F0FA"}

def gtag(game):
    return {"IG":"Jeu · Instant gagnant","AUTO":"Jeu · 100 % gagnant","TAS":"Jeu · Tirage au sort"}[game]

def make_tests(game,oblig):
    t=[("Accès & états","Accéder à l'URL de test","La page se charge, la barre Houston est visible en haut"),
       ("Accès & états","État « En attente »","La page d'attente s'affiche (jeu pas encore commencé)"),
       ("Accès & états","État « En cours »","Le bouton « Je participe / Je joue » apparaît, le parcours est accessible"),
       ("Accès & états","État « Terminé »","La page de fin de jeu s'affiche"),
       ("Accès & états","État « Offline »","La page « opération suspendue » s'affiche"),
       ("Accès & états","État « Quota atteint » (si compteur)","La page « limite de participations atteinte » s'affiche"),
       ("Home page","Affichage Home desktop","Visuels, logos, textes et charte conformes à la maquette validée"),
       ("Home page","Affichage Home mobile (responsive)","Mise en page correcte sur smartphone, aucun élément coupé"),
       ("Home page","Mentions légales / règlement du jeu / confidentialité","Liens présents et documents accessibles"),
       ("Home page","Bouton « Je participe / Je joue »","Lance le parcours de participation"),
       ("Formulaire","Présence de tous les champs attendus","Tous les champs du cahier des charges sont présents"),
       ("Formulaire","Champs obligatoires laissés vides","Message d'erreur clair, blocage de la validation"),
       ("Formulaire","Contrôles de format (e-mail, code postal, tél.)","Une saisie invalide déclenche un message d'erreur")]
    if oblig=="AOA":
        t+=[("Preuve d'achat","Upload d'un justificatif valide (PDF/JPG/PNG)","Le fichier est accepté et visible dans le récapitulatif"),
            ("Preuve d'achat","Upload d'un format / poids non autorisé","Message d'erreur, fichier refusé"),
            ("Preuve d'achat","Validation sans justificatif","Blocage tant que le justificatif obligatoire n'est pas fourni")]
    else:
        t+=[("Participation SOA","Aucune preuve d'achat demandée","Le parcours ne réclame aucun justificatif d'achat"),
            ("Participation SOA","Mention SOA / remboursement des frais","La participation gratuite et le remboursement des frais sont indiqués (règlement)")]
    if game=="IG":
        t+=[("Jeu · Instant gagnant","Forcer GAGNE (via Houston) et jouer","Page « Vous avez gagné » + animation + lot affiché"),
            ("Jeu · Instant gagnant","Forcer PERD (via Houston) et jouer","Page « Vous avez perdu » cohérente (relance / remerciement)"),
            ("Jeu · Instant gagnant","Rejouer après avoir déjà joué","Message « vous avez déjà participé » selon la règle")]
    elif game=="AUTO":
        t+=[("Jeu · Gros lot (IG)","Forcer GAGNE (grille un instant gagnant) et jouer","Page « gagné » + gros lot affiché + animation"),
            ("Jeu · Lot de consolation","Forcer PERD (pas d'IG grillé) et jouer","Page « lot de consolation » (ex. bon de réduction) — tout le monde repart avec un lot, pas de page perdant sèche"),
            ("Jeu · 100 % gagnant","Rejouer après avoir déjà joué","Message « vous avez déjà participé » selon la règle")]
    else:
        t+=[("Jeu · Tirage au sort","Valider une participation","Message « participation enregistrée, tirage le [date] », aucun résultat immédiat"),
            ("Jeu · Tirage au sort","Absence de résultat immédiat","Le parcours ne révèle pas de gagné/perdu (désignation au tirage)")]
    if game=="TAS":
        t+=[("Dotation / gain","Dotation et date de tirage (règlement)","Le règlement précise la dotation et la date de tirage"),
            ("Dotation / gain","Process e-mail après tirage","L'envoi de l'e-mail aux gagnants après tirage est prévu (à contrôler au tirage)")]
    elif game=="AUTO":
        t+=[("Dotation / gain","Attribution du bon lot selon le résultat","Gros lot si instant gagnant grillé, lot de consolation sinon — le bon lot s'affiche"),
            ("Dotation / gain","E-mail de gain","E-mail reçu avec le lot correspondant (gros lot ou bon de réduction) + instructions")]
    else:
        t+=[("Dotation / gain","Affichage du lot gagné","Le bon lot s'affiche (visuel, valeur, code si dématérialisé)"),
            ("Dotation / gain","E-mail de gain","E-mail reçu avec le lot / code et les instructions pour en profiter")]
    t+=[("Consentements","Case CGU / règlement / RGPD non cochée","Impossible de valider tant que le consentement obligatoire n'est pas donné"),
        ("Anti-fraude","Dédoublonnage (contrôles activés)","Rejeu bloqué selon la règle (ex. 1 participation / jour) avec message"),
        ("Anti-fraude","Désactiver contrôles","Permet de rejouer avec le même e-mail pour tester"),
        ("E-mails","E-mail de confirmation de participation","E-mail reçu ; expéditeur, objet, visuels et liens corrects"),
        ("E-mails","Rendu e-mail sur mobile","E-mail lisible et bien affiché sur smartphone")]
    # Pas de suivi de participation pour un jeu SANS obligation d'achat (SOA)
    if oblig!="SOA":
        t.append(("E-mails","Suivi de participation (pied de page)","Après saisie de l'e-mail, réception du lien de suivi"))
    return t

def build(game,oblig):
    gm=GAME_META[game]; om=OBL_META[oblig]; AOA=(oblig=="AOA")
    wb=openpyxl.Workbook()
    # ---------- 1) MODE D'EMPLOI ----------
    ws=wb.active; ws.title="Mode d'emploi"; ws.sheet_view.showGridLines=False
    ws.column_dimensions["A"].width=2.5; ws.column_dimensions["B"].width=40
    ws.column_dimensions["C"].width=76; ws.column_dimensions["D"].width=3
    r=title_bar(ws,2,3,2,f"PROCÉDURE DE TEST — JEU {gm['name']}",
                f"{om['name'].capitalize()} · gabarit générique — à dupliquer et compléter pour chaque opération"); r+=1
    hc=ws.cell(r,2,"  À LIRE EN PREMIER"); hc.font=F(12,True,WHITE); hc.fill=fill(ACCENT)
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3); hc.alignment=leftc; ws.row_dimensions[r].height=22; r+=1
    intro=("Vous allez tester le site de votre jeu-concours avant sa mise en ligne. "
           "Ce site « gabarit » est une copie de test, identique à ce que verront les participants. "
           f"Objectif : vérifier le parcours, le formulaire, le mécanisme de jeu ({gm['mech']}) et les e-mails, puis valider avant le lancement.\n\n"
           "Comptez environ 30 minutes. Avant de commencer, munissez-vous d'une boîte mail que vous pouvez consulter (pensez à vérifier les spams)"
           +(", et d'un fichier à téléverser en guise de justificatif d'achat" if AOA else "")+
           ". Testez de préférence sur ordinateur puis sur mobile. En cas de blocage, contactez votre chef de projet Promo.dev (coordonnées ci-dessous).")
    ic=ws.cell(r,2,intro); ic.font=F(11); ic.alignment=Alignment(wrap_text=True,vertical="top",indent=1); ic.fill=fill("FBF3E6")
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    for cc in (2,3): ws.cell(r,cc).border=border
    ws.row_dimensions[r].height=150; r+=2
    section(ws,2,r,"1 · Informations de l'opération"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    tag=ws.cell(r,2,"▸ À compléter par le chef de projet Promo.dev avant l'envoi au client")
    tag.font=F(10,True,ACCENT); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    tag.alignment=Alignment(indent=1,vertical="center"); ws.row_dimensions[r].height=20; r+=1
    info=[("Client / Marque","ex. Marque X"),("Nom de l'opération","ex. Jeu Rentrée 2026"),
     ("N° opération (idgame)","ex. 999"),("URL de test","ex. https://test.promo.dev/xxxxxxxx"),
     ("Type de jeu",gm['type_val']),
     ("Obligation d'achat","Oui — justificatif requis" if AOA else "Non — participation SOA")]
    if AOA: info.append(("Justificatif attendu","ex. ticket de caisse / facture (PDF, JPG, PNG)"))
    info+=[("Dotation / lots","ex. 100 cartes cadeaux 50 €"+(" + bons de réduction (consolation)" if game=="AUTO" else "")),
     ("E-mail de test","ex. prenom.nom+test@domaine.fr"),
     ("Dédoublonnage paramétré","ex. 1 participation / jour / e-mail"),
     ("Dates (début / fin"+(" / tirage" if gm['tirage'] else "")+")","ex. 30/09/2026 → 31/01/2027"+(", tirage le 15/02/2027" if gm['tirage'] else "")),
     ("Testeur(s) côté client","ex. Prénom Nom"),("Contact Promo.dev (chef de projet)","ex. prenom@promo.dev"),
     ("Date limite de retour des tests","ex. JJ/MM/AAAA")]
    for label,exv in info:
        lc=ws.cell(r,2,label); lc.font=F(11,True,"26324A"); lc.fill=fill(GREY); lc.alignment=leftc; bd(ws,lc.coordinate)
        val=INFO.get(label,exv); is_real=(label in INFO) or label in ("Type de jeu","Obligation d'achat")
        vc=ws.cell(r,3,val); vc.font=F(11,False,"26324A" if is_real else "9AA0AA"); vc.fill=fill(YELLOW); vc.alignment=leftc
        yb=Side(style="thin",color=YELLOWB); vc.border=Border(left=yb,right=yb,top=yb,bottom=yb)
        ws.row_dimensions[r].height=26; r+=1
    r+=1
    section(ws,2,r,"2 · Comment utiliser ce classeur"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    steps=["Vérifiez les informations de l'opération ci-dessus (pré-remplies par Promo.dev) : URL, type de jeu, dotation, dates.",
     "Onglet « Guide Houston » : familiarisez-vous avec la barre d'administration et les états du jeu.",
     "Onglet « Checklist » : déroulez chaque test, indiquez un statut par appareil (OK / KO / N.A.) et un commentaire.",
     "Onglet « Anomalies » : décrivez précisément chaque bug (page, action, capture).",
     "Onglet « Validation » : donnez votre bon pour mise en ligne une fois tous les tests OK.",
     "Renvoyez le fichier complété au chef de projet avant la date limite."]
    for i,s in enumerate(steps,1):
        b=ws.cell(r,2,f"Étape {i}"); b.font=F(11,True,ACCENT); b.alignment=Alignment(vertical="center",indent=1)
        b.fill=fill(ZEBRA if i%2==0 else WHITE)
        c=ws.cell(r,3,s); c.font=F(11); c.alignment=wrapv; c.fill=fill(ZEBRA if i%2==0 else WHITE)
        ws.row_dimensions[r].height=32; r+=1
    r+=1
    section(ws,2,r,"3 · Légende des statuts"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    for code,desc,bg,tx in [("OK","Conforme, rien à signaler",GREENF,GREENT),
     ("KO","Anomalie détectée → à détailler dans l'onglet Anomalies",REDF,REDT),
     ("N.A.","Non applicable à cette opération",NEUTF,"555555"),("À tester","Test pas encore réalisé",YELLOW,"B45309")]:
        p=ws.cell(r,2,code); p.font=F(11,True,tx); p.fill=fill(bg); p.alignment=center; bd(ws,p.coordinate)
        d=ws.cell(r,3,desc); d.font=F(11); d.alignment=leftc; ws.row_dimensions[r].height=22; r+=1
    r+=1
    section(ws,2,r,"4 · Lexique"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    lex=[("Jeu-concours","Opération où le participant tente de gagner un lot, en ligne."),
     ("Site gabarit","Version de test du site de l'opération, identique au site final, pour tout valider avant la mise en ligne."),
     ("Houston","Outil interne Promo.dev (la barre en haut du site de test) qui sert à simuler les états et le résultat du jeu. Invisible pour les participants."),
     ("Règlement du jeu","Document légal encadrant le jeu (conditions, dotation, dates, tirage)."),
     ("Dotation / lot","Ensemble des cadeaux mis en jeu.")]
    if game=="IG": lex.append(("Instant gagnant","Le résultat gagné / perdu s'affiche immédiatement après la participation."))
    if game=="AUTO": lex+=[("100 % gagnant","Jeu instant gagnant où tout le monde repart avec un lot : gros lot pour ceux qui grillent un instant gagnant, lot de consolation (ex. bon de réduction) pour les autres."),
                           ("Lot de consolation","Petit lot (ex. bon de réduction) attribué aux participants qui ne remportent pas le gros lot.")]
    if game=="TAS": lex.append(("Tirage au sort","Les gagnants sont désignés par tirage à une date donnée ; pas de résultat immédiat."))
    if AOA: lex.append(("Justificatif / preuve d'achat","Preuve d'achat demandée au participant (ticket de caisse, facture)."))
    else: lex.append(("SOA — Sans obligation d'achat","Voie de participation gratuite, sans achat, imposée par la loi pour les jeux."))
    lex.append(("Dédoublonnage","Contrôle qui empêche les participations en double (ex. 1 par jour / e-mail)."))
    lex.append(("Bon pour mise en ligne","Validation finale du client attestant que l'opération peut être publiée."))
    for term,desc in lex:
        tcell=ws.cell(r,2,term); tcell.font=F(11,True,"26324A"); tcell.fill=fill(GREY)
        tcell.alignment=Alignment(vertical="top",wrap_text=True,indent=1); bd(ws,tcell.coordinate)
        dcell=ws.cell(r,3,desc); dcell.font=F(11); dcell.alignment=Alignment(vertical="top",wrap_text=True,indent=1); bd(ws,dcell.coordinate)
        ws.row_dimensions[r].height=32; r+=1
    ws.page_setup.orientation="portrait"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ws.print_area=f"A1:D{r}"
    build_guide(wb,game,oblig); build_checklist(wb,game,oblig); build_anomalies(wb); build_validation(wb)
    return wb

def build_guide(wb,game,oblig):
    gm=GAME_META[game]
    g=wb.create_sheet("Guide Houston"); g.sheet_view.showGridLines=False
    g.column_dimensions["A"].width=2.5; g.column_dimensions["B"].width=30
    g.column_dimensions["C"].width=60; g.column_dimensions["D"].width=60
    r=title_bar(g,2,4,2,"GUIDE DE LA BARRE D'ADMINISTRATION « HOUSTON »",
                "La barre en haut du site de test permet de simuler les statuts, pages et résultats du jeu."); r+=1
    hh=g.cell(r,2,"⚠  La barre « Houston » est un outil de test interne"); hh.font=F(12,True,WHITE); hh.fill=fill(NAVY)
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); hh.alignment=leftc; g.row_dimensions[r].height=24; r+=1
    frm=("Elle apparaît uniquement sur ce site de test et sert à simuler les états du jeu "
         "(en cours, en attente, terminé…)"+(" et à forcer le résultat du jeu (gagné / perdu)" if game in("IG","AUTO") else "")+
         ". Elle ne sera JAMAIS visible par les participants sur le site final.")
    fc=g.cell(r,2,frm); fc.font=F(11); fc.alignment=Alignment(wrap_text=True,vertical="center",indent=1); fc.fill=fill("FBF3E6")
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    for cc in range(2,5): g.cell(r,cc).border=border
    g.row_dimensions[r].height=48; r+=2
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    for cc in range(2,5): g.cell(r,cc).fill=fill(TINT); g.cell(r,cc).border=border
    place_abs(g,"image9.png",2,r,720,90,3,90); r+=2
    qh=g.cell(r,2,"  PAR OÙ COMMENCER — testez en 6 étapes"); qh.font=F(12,True,WHITE); qh.fill=fill(ACCENT)
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); qh.alignment=leftc; g.row_dimensions[r].height=22; r+=1
    if game=="IG": s4="Forcez le résultat via « Participation » : GAGNE puis PERD, pour tester la page gagné ET la page perdu."
    elif game=="AUTO": s4="Forcez « Participation » : GAGNE (gros lot) puis PERD (lot de consolation, ex. bon de réduction) — tout le monde repart avec un lot."
    else: s4="Validez la participation et vérifiez le message « participation enregistrée, tirage le [date] » (pas de résultat immédiat)."
    s5=("Vérifiez la page résultat (lot affiché) et les e-mails reçus." if game in("IG","AUTO")
        else "Vérifiez le message de participation et l'e-mail de confirmation.")
    qs=["Ouvrez l'URL de test (voir l'onglet « Mode d'emploi »).",
        "Dans la barre Houston, réglez « État » sur « En cours ».",
        "Cliquez sur « Je participe / Je joue » et remplissez le formulaire.",
        s4, s5,
        "Notez vos remarques dans « Checklist » (OK/KO) et détaillez tout souci dans « Anomalies »."]
    for i,s in enumerate(qs,1):
        zeb=ZEBRA if i%2==0 else WHITE
        eb=g.cell(r,2,f"Étape {i}"); eb.font=F(11,True,ACCENT); eb.alignment=Alignment(vertical="top",indent=1); eb.fill=fill(zeb); bd(g,eb.coordinate)
        sc=g.cell(r,3,s); sc.font=F(11); sc.alignment=wrap; sc.fill=fill(zeb)
        g.merge_cells(start_row=r,start_column=3,end_row=r,end_column=4); g.cell(r,4).fill=fill(zeb)
        bd(g,f"C{r}"); bd(g,f"D{r}")
        g.row_dimensions[r].height=42 if i!=4 else 56; r+=1
    r+=1
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    nc=g.cell(r,2,"ℹ Les aperçus ci-dessous sont des exemples génériques. La barre d'administration Houston est identique pour toutes les opérations ; seules les pages du jeu changent selon la marque.")
    nc.font=F(10,True,"6B7280"); nc.fill=fill("F2F4F8"); nc.alignment=Alignment(wrap_text=True,vertical="center",indent=1)
    for cc in range(2,5): g.cell(r,cc).border=border
    g.row_dimensions[r].height=32; r+=2
    hr=r
    for i,txt in enumerate(("CONTRÔLE","À QUOI ÇA SERT","APERÇU")):
        c=g.cell(hr,2+i,txt); c.font=F(11,True,WHITE); c.fill=fill(BLUE); c.alignment=wrapc; bd(g,c.coordinate)
    g.row_dimensions[hr].height=24
    # ligne "visuel" adaptée : pages du jeu (pas de "page done")
    if game=="IG":
        vrow=("VISUEL","Permet de visualiser directement les pages du jeu : la Home, la page « Perdu », et — pour chaque lot mis en jeu — sa page « Gagné » (le menu liste le nom de chaque lot, ex. « e-carte 20 € »).","visuel_ig.png",400)
    elif game=="AUTO":
        vrow=("VISUEL","Permet de visualiser directement les pages du jeu : la Home, la page « Lot de consolation », et — pour chaque gros lot — sa page « Gagné » (le menu liste le nom de chaque lot).","visuel_auto.png",400)
    else:
        vrow=("VISUEL","Permet de visualiser directement la Home et la page de confirmation de participation.","visuel_tas.png",400)
    # ligne "participation" adaptée
    if game=="IG":
        prow=("PARTICIPATION (résultat) ★","CLÉ POUR UN JEU INSTANT GAGNANT : force le résultat GAGNE ou PERD pour tester la page « Vous avez gagné » (avec le lot) ET la page « Vous avez perdu ». À tester dans les deux positions.","image8.png",240)
    elif game=="AUTO":
        prow=("PARTICIPATION (résultat) ★","CLÉ POUR CE JEU 100 % GAGNANT : GAGNE = le participant grille un instant gagnant (gros lot) ; PERD = pas d'IG grillé mais un lot de consolation (ex. bon de réduction). Tester les deux pages.","image8.png",240)
    else:
        prow=("PARTICIPATION (résultat)","Tirage au sort : le résultat n'est pas immédiat (désignation au tirage). Ce sélecteur n'a pas d'effet visible sur le parcours participant.","image8.png",240)
    guide=[("CONNEXION","Se connecter sur l'URL de test de l'opération (voir onglet « Mode d'emploi »).",None,0),
     ("ÉTAT","Permet de tester les différents statuts de l'opération (menu déroulant).","image6.png",300),
     ("ÉTAT ▸ En cours","Le parcours est accessible ; le bouton « Je participe / Je joue » apparaît pour lancer et valider une participation.","image1.png",210),
     ("ÉTAT ▸ En attente","Affiche la page d'attente : le jeu n'a pas encore commencé.","neutral_attente_jeu.png",380),
     ("ÉTAT ▸ Terminé","Affiche la page de fin : le jeu est terminé.","neutral_termine_jeu.png",380),
     ("ÉTAT ▸ Offline","Affiche la page « hors ligne » quand le jeu est suspendu.","neutral_offline_jeu.png",380),
     ("ÉTAT ▸ Quota atteint","Affiche la page d'accueil quand le nombre max de participations est atteint (si compteur limitatif).","neutral_quota_jeu.png",380),
     vrow,
     prow,
     ("DÉSACTIVER CONTRÔLES","Désactive le contrôle d'unicité : permet de rejouer plusieurs fois avec le même e-mail. Décoché = le dédoublonnage paramétré s'applique (ex. 1 participation / jour).","image3.png",260)]
    # Pas de suivi de participation pour un jeu SANS obligation d'achat (SOA)
    if oblig!="SOA":
        guide.append(("SUIVI DE PARTICIPATION","Depuis le pied de page, permet de recevoir un e-mail avec le lien de suivi de la participation.","neutral_suivi.png",360))
    row=hr+1
    for idx,(label,desc,imgn,tw) in enumerate(guide):
        zeb=ZEBRA if idx%2==0 else WHITE
        lc=g.cell(row,2,label); lc.font=F(11,True,NAVY); lc.alignment=wrapv; lc.fill=fill(TINT); bd(g,lc.coordinate)
        dc=g.cell(row,3,desc); dc.font=F(12); dc.alignment=wrap; dc.fill=fill(zeb); bd(g,dc.coordinate)
        ac=g.cell(row,4); ac.fill=fill(zeb); bd(g,ac.coordinate)
        if imgn:
            place_abs(g,imgn,4,row,tw,230,1,60)
        else:
            g.row_dimensions[row].height=42
        row+=1
    g.page_setup.orientation="landscape"; g.page_setup.fitToWidth=1; g.page_setup.fitToHeight=0
    g.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); g.print_area=f"A1:D{row}"

def build_checklist(wb,game,oblig):
    gm=GAME_META[game]; om=OBL_META[oblig]
    ck=wb.create_sheet("Checklist"); ck.sheet_view.showGridLines=False
    for col,w in {"A":2.5,"B":5.5,"C":23,"D":46,"E":46,"F":12,"G":12,"H":12,"I":34}.items(): ck.column_dimensions[col].width=w
    r=title_bar(ck,2,9,2,f"CHECKLIST — JEU {gm['name']} ({om['name']})",
                "Renseignez un statut par appareil (Desktop / Tablette / Mobile). Tout KO doit être détaillé dans l'onglet « Anomalies »."); r+=1
    headers=["N°","Catégorie","Test à réaliser","Résultat attendu","Desktop","Tablette","Mobile","Commentaire / Anomalie"]
    hr=r
    for i,h in enumerate(headers):
        c=ck.cell(hr,2+i,h); c.font=F(11,True,WHITE); c.fill=fill(BLUE); c.alignment=wrapc; bd(ck,c.coordinate)
    ck.row_dimensions[hr].height=32
    tests=make_tests(game,oblig)
    row=hr+1
    exg={"IG":("Jeu · Instant gagnant","Forcer GAGNE via Houston et jouer","Page « Vous avez gagné » + lot affiché"),
         "AUTO":("Jeu · Lot de consolation","Forcer PERD (pas d'IG) et jouer","Page « lot de consolation » (bon de réduction)"),
         "TAS":("Jeu · Tirage au sort","Valider une participation","Message « participation enregistrée, tirage le [date] »")}[game]
    ex=("Ex",exg[0],exg[1],exg[2],"OK","OK","KO","KO mobile : affichage")
    for i,v in enumerate(ex):
        c=ck.cell(row,2+i,v); c.border=border; c.alignment=(wrapc if i in(0,4,5,6) else wrap); c.font=F(10,it=True,color="9AA0AA"); c.fill=fill("FBFBFD")
    ck.row_dimensions[row].height=28; row+=1
    start=row; n=1
    for cat,test,exp in tests:
        zeb=ZEBRA if n%2==0 else WHITE
        a=ck.cell(row,2,n); a.alignment=center; a.font=F(11,True,"6B7280"); a.fill=fill(zeb); bd(ck,a.coordinate)
        cc=ck.cell(row,3,cat); cc.font=F(11,True,"2C3E63"); cc.alignment=wrapv; cc.fill=fill(CATCOL.get(cat,"EDEFF3")); bd(ck,cc.coordinate)
        t=ck.cell(row,4,test); t.font=F(11); t.alignment=wrap; t.fill=fill(zeb); bd(ck,t.coordinate)
        e=ck.cell(row,5,exp); e.font=F(11,color="4B5563"); e.alignment=wrap; e.fill=fill(zeb); bd(ck,e.coordinate)
        for col in (6,7,8):
            x=ck.cell(row,col); x.fill=fill(zeb); x.alignment=center; bd(ck,x.coordinate); x.font=F(11)
        cm=ck.cell(row,9); cm.fill=fill(zeb); cm.alignment=wrap; bd(ck,cm.coordinate); cm.font=F(11)
        ck.row_dimensions[row].height=34; row+=1; n+=1
    end=row-1
    dv=DataValidation(type="list",formula1='"OK,KO,N.A.,À tester"',allow_blank=True); dv.add(f"F{start}:H{end}"); ck.add_data_validation(dv)
    for val,bg,tx in [("OK",GREENF,GREENT),("KO",REDF,REDT),("N.A.",NEUTF,"555555"),("À tester",YELLOW,"B45309")]:
        ck.conditional_formatting.add(f"F{start}:H{end}",CellIsRule(operator="equal",formula=[f'"{val}"'],fill=fill(bg),font=F(11,True,tx)))
    sr=end+2
    ck.merge_cells(start_row=sr,start_column=3,end_row=sr,end_column=9)
    hh=ck.cell(sr,3,"  RÉCAPITULATIF (nombre de tests par statut et par appareil)"); hh.font=F(12,True,WHITE); hh.fill=fill(NAVY); hh.alignment=leftc; ck.row_dimensions[sr].height=24
    ck.merge_cells(start_row=sr+1,start_column=3,end_row=sr+1,end_column=5)
    for cidx,dev in [(6,"Desktop"),(7,"Tablette"),(8,"Mobile")]:
        dc=ck.cell(sr+1,cidx,dev); dc.font=F(11,True,NAVY); dc.fill=fill(TINT); dc.alignment=center; bd(ck,dc.coordinate)
    rr=sr+2
    for lab,kind,bg,tx in [("Tests OK","OK",GREENF,GREENT),("Tests KO","KO",REDF,REDT),("Restant à tester","REST",YELLOW,"B45309")]:
        ck.merge_cells(start_row=rr,start_column=3,end_row=rr,end_column=5)
        l=ck.cell(rr,3,lab); l.font=F(11,True,tx); l.fill=fill(bg); l.alignment=leftc; bd(ck,l.coordinate)
        for cidx,cl in [(6,"F"),(7,"G"),(8,"H")]:
            if kind=="REST":
                form=f'=COUNTA(D{start}:D{end})-COUNTIF({cl}{start}:{cl}{end},"OK")-COUNTIF({cl}{start}:{cl}{end},"KO")-COUNTIF({cl}{start}:{cl}{end},"N.A.")'
            else:
                form=f'=COUNTIF({cl}{start}:{cl}{end},"{kind}")'
            v=ck.cell(rr,cidx,form); v.font=F(12,True,tx); v.fill=fill(bg); v.alignment=center; bd(ck,v.coordinate)
        ck.row_dimensions[rr].height=22; rr+=1
    ck.freeze_panes=f"B{hr+1}"
    ck.page_setup.orientation="landscape"; ck.page_setup.fitToWidth=1; ck.page_setup.fitToHeight=0
    ck.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ck.print_area=f"A1:I{rr}"

def build_anomalies(wb):
    an=wb.create_sheet("Anomalies"); an.sheet_view.showGridLines=False
    for c,w in {"A":2.5,"B":5.5,"C":14,"D":24,"E":42,"F":36,"G":14,"H":14,"I":24}.items(): an.column_dimensions[c].width=w
    r=title_bar(an,2,9,2,"JOURNAL DES ANOMALIES",
                "Une ligne par bug. Reliez chaque anomalie au N° de test concerné et joignez une capture si possible."); r+=1
    hdr=["N°","N° test lié","Page / Écran","Description du problème","Étapes pour reproduire","Gravité","Statut","Capture (nom fichier)"]
    hr=r
    for i,h in enumerate(hdr):
        c=an.cell(hr,2+i,h); c.font=F(11,True,WHITE); c.fill=fill(BLUE); c.alignment=wrapc; bd(an,c.coordinate)
    an.row_dimensions[hr].height=32
    exa=("Ex","16","Page résultat","L'animation « gagné » reste figée sur mobile","1. Forcer GAGNE  2. Jouer sur smartphone","Majeure","Ouvert","capture_01.png")
    row=hr+1
    for i,v in enumerate(exa):
        c=an.cell(row,2+i,v); c.border=border; c.alignment=(wrapc if i in(0,1,5,6) else wrap); c.font=F(10,it=True,color="9AA0AA"); c.fill=fill("FBFBFD")
    an.row_dimensions[row].height=30; row+=1
    first=row
    for k in range(20):
        zeb=ZEBRA if k%2==0 else WHITE
        for col in range(2,10):
            c=an.cell(row,col); c.border=border; c.font=F(11); c.fill=fill(zeb); c.alignment=(wrapc if col in(2,3,7,8) else wrap)
        an.row_dimensions[row].height=28; row+=1
    last=row-1
    dvg=DataValidation(type="list",formula1='"Bloquante,Majeure,Mineure,Cosmétique"',allow_blank=True); dvg.add(f"G{first}:G{last}"); an.add_data_validation(dvg)
    dvs=DataValidation(type="list",formula1='"Ouvert,Corrigé,Vérifié,Rejeté"',allow_blank=True); dvs.add(f"H{first}:H{last}"); an.add_data_validation(dvs)
    an.conditional_formatting.add(f"G{first}:G{last}",CellIsRule(operator="equal",formula=['"Bloquante"'],fill=fill(REDF),font=F(11,True,REDT)))
    an.conditional_formatting.add(f"G{first}:G{last}",CellIsRule(operator="equal",formula=['"Majeure"'],fill=fill("FCE5CD"),font=F(11,True,"B45309")))
    an.conditional_formatting.add(f"H{first}:H{last}",CellIsRule(operator="equal",formula=['"Corrigé"'],fill=fill(GREENF),font=F(11,True,GREENT)))
    an.conditional_formatting.add(f"H{first}:H{last}",CellIsRule(operator="equal",formula=['"Vérifié"'],fill=fill(GREENF),font=F(11,True,GREENT)))
    an.freeze_panes=f"B{hr+1}"
    an.page_setup.orientation="landscape"; an.page_setup.fitToWidth=1; an.page_setup.fitToHeight=0
    an.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); an.print_area=f"A1:I{last}"

def build_validation(wb):
    va=wb.create_sheet("Validation"); va.sheet_view.showGridLines=False
    va.column_dimensions["A"].width=2.5; va.column_dimensions["B"].width=44; va.column_dimensions["C"].width=54; va.column_dimensions["D"].width=3
    r=title_bar(va,2,3,2,"VALIDATION — BON POUR MISE EN LIGNE",
                "À compléter une fois la checklist terminée. Cette validation vaut accord pour la mise en ligne de l'opération."); r+=1
    valrow=None
    for label,val in [("Opération validée","__DROPDOWN__"),
     ("Réserves / points en suspens",""),("Nombre d'anomalies bloquantes restantes",""),
     ("Nom du valideur (client)",""),("Fonction",""),("Date de validation",""),
     ("Signature (nom + « bon pour accord »)",""),("Chef de projet Promo.dev","")]:
        l=va.cell(r,2,label); l.font=F(11,True,"26324A"); l.fill=fill(GREY); l.alignment=leftc; bd(va,l.coordinate)
        c=va.cell(r,3); c.fill=fill(YELLOW); c.alignment=Alignment(vertical="center",wrap_text=True,indent=1)
        yb=Side(style="thin",color=YELLOWB); c.border=Border(left=yb,right=yb,top=yb,bottom=yb)
        if val=="__DROPDOWN__":
            c.value="— cliquez pour choisir —"; c.font=F(11,it=True,color="9AA0AA"); valrow=r
        else:
            c.value=val; c.font=F(11)
        va.row_dimensions[r].height=38 if val=="" else 28; r+=1
    dvv=DataValidation(type="list",formula1='"Oui sans réserve,Oui avec réserves,Non"',allow_blank=True)
    dvv.add(f"C{valrow}"); va.add_data_validation(dvv)
    va.conditional_formatting.add(f"C{valrow}",CellIsRule(operator="equal",formula=['"Oui sans réserve"'],fill=fill(GREENF),font=F(11,True,GREENT)))
    va.conditional_formatting.add(f"C{valrow}",CellIsRule(operator="equal",formula=['"Oui avec réserves"'],fill=fill(YELLOW),font=F(11,True,"B45309")))
    va.conditional_formatting.add(f"C{valrow}",CellIsRule(operator="equal",formula=['"Non"'],fill=fill(REDF),font=F(11,True,REDT)))
    va.page_setup.orientation="portrait"; va.page_setup.fitToWidth=1; va.page_setup.fitToHeight=0
    va.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); va.print_area=f"A1:D{r}"


def build_prime(wb):


    # ================= 1) MODE D'EMPLOI =================
    ws=wb.active; ws.title="Mode d'emploi"; ws.sheet_view.showGridLines=False
    ws.column_dimensions["A"].width=2.5; ws.column_dimensions["B"].width=40
    ws.column_dimensions["C"].width=76; ws.column_dimensions["D"].width=3
    r=title_bar(ws,2,3,2,"PROCÉDURE DE TEST — OFFRE DE PRIME (PRIME)",
                "Gabarit générique · à dupliquer et compléter pour chaque opération"); r+=1
    hc=ws.cell(r,2,"  À LIRE EN PREMIER"); hc.font=F(12,True,WHITE); hc.fill=fill(ACCENT)
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3); hc.alignment=leftc; ws.row_dimensions[r].height=22; r+=1
    intro=("Vous allez tester le site de votre offre de prime avant sa mise en ligne. Une prime, c'est un cadeau "
           "(souvent un produit) offert en contrepartie d'un achat. Ce site « gabarit » est une copie de test, identique à ce que verront les participants. "
           "Objectif : vérifier le parcours, le formulaire, l'envoi du justificatif d'achat et de l'adresse de livraison, et les e-mails, puis valider avant le lancement.\n\n"
           "Comptez environ 30 minutes. Avant de commencer, munissez-vous : d'une boîte mail que vous pouvez consulter (pensez à vérifier les spams), "
           "d'un fichier à téléverser en guise de justificatif d'achat (facture ou ticket), et d'une adresse de test. "
           "Testez de préférence sur ordinateur puis sur mobile. En cas de blocage, contactez votre chef de projet Promo.dev (coordonnées ci-dessous).")
    ic=ws.cell(r,2,intro); ic.font=F(11); ic.alignment=Alignment(wrap_text=True,vertical="top",indent=1); ic.fill=fill("FBF3E6")
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    for cc in (2,3): ws.cell(r,cc).border=border
    ws.row_dimensions[r].height=160; r+=2
    section(ws,2,r,"1 · Informations de l'opération"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    tag=ws.cell(r,2,"▸ À compléter par le chef de projet Promo.dev avant l'envoi au client")
    tag.font=F(10,True,ACCENT); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    tag.alignment=Alignment(indent=1,vertical="center"); ws.row_dimensions[r].height=20; r+=1
    info=[("Client / Marque","ex. Marque (société)"),("Nom de l'opération","ex. Offre de lancement smartphone"),
     ("N° opération (idgame)","ex. 999"),("URL de test","ex. https://test.promo.dev/xxxxxxxx"),
     ("Nature de la prime","ex. produit offert (montre connectée)"),
     ("Justificatif attendu","ex. facture / ticket de caisse (PDF, JPG, PNG)"),
     ("Adresse de livraison requise","Oui (le cadeau est expédié au participant)"),
     ("Preuve produit demandée ?","ex. n° de série / IMEI  —  ou : non"),
     ("E-mail de test","ex. prenom.nom+test@domaine.fr"),
     ("Dédoublonnage paramétré","ex. 1 prime / foyer / justificatif"),
     ("Dates (début / fin d'achat / fin)","ex. 05/05/2026 → 30/06/2026"),
     ("Testeur(s) côté client","ex. Prénom Nom"),("Contact Promo.dev (chef de projet)","ex. prenom@promo.dev"),
     ("Date limite de retour des tests","ex. JJ/MM/AAAA")]
    for label,exv in info:
        lc=ws.cell(r,2,label); lc.font=F(11,True,"26324A"); lc.fill=fill(GREY); lc.alignment=leftc; bd(ws,lc.coordinate)
        val=INFO.get(label,exv); is_real=(label in INFO) or label in ("Adresse de livraison requise",)
        vc=ws.cell(r,3,val); vc.font=F(11,False,"26324A" if is_real else "9AA0AA"); vc.fill=fill(YELLOW); vc.alignment=leftc
        yb=Side(style="thin",color=YELLOWB); vc.border=Border(left=yb,right=yb,top=yb,bottom=yb)
        ws.row_dimensions[r].height=26; r+=1
    r+=1
    section(ws,2,r,"2 · Comment utiliser ce classeur"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    steps=["Vérifiez les informations de l'opération ci-dessus (pré-remplies par Promo.dev) : URL, nature de la prime, justificatif, dates.",
     "Onglet « Guide Houston » : familiarisez-vous avec la barre d'administration et les états de l'offre.",
     "Onglet « Checklist PRIME » : déroulez chaque test, indiquez un statut par appareil (OK / KO / N.A.) et un commentaire.",
     "Onglet « Anomalies » : décrivez précisément chaque bug (page, action, capture).",
     "Onglet « Validation » : donnez votre bon pour mise en ligne une fois tous les tests OK.",
     "Renvoyez le fichier complété au chef de projet avant la date limite."]
    for i,s in enumerate(steps,1):
        b=ws.cell(r,2,f"Étape {i}"); b.font=F(11,True,ACCENT); b.alignment=Alignment(vertical="center",indent=1)
        b.fill=fill(ZEBRA if i%2==0 else WHITE)
        c=ws.cell(r,3,s); c.font=F(11); c.alignment=wrapv; c.fill=fill(ZEBRA if i%2==0 else WHITE)
        ws.row_dimensions[r].height=32; r+=1
    r+=1
    section(ws,2,r,"3 · Légende des statuts"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    for code,desc,bg,tx in [("OK","Conforme, rien à signaler",GREENF,GREENT),
     ("KO","Anomalie détectée → à détailler dans l'onglet Anomalies",REDF,REDT),
     ("N.A.","Non applicable à cette opération",NEUTF,"555555"),("À tester","Test pas encore réalisé",YELLOW,"B45309")]:
        p=ws.cell(r,2,code); p.font=F(11,True,tx); p.fill=fill(bg); p.alignment=center; bd(ws,p.coordinate)
        d=ws.cell(r,3,desc); d.font=F(11); d.alignment=leftc; ws.row_dimensions[r].height=22; r+=1
    r+=1
    section(ws,2,r,"4 · Lexique"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    lex=[("Offre de prime (PRIME)","Opération où le consommateur reçoit une prime (un cadeau, souvent un produit) en contrepartie d'un achat, sur présentation d'un justificatif."),
     ("Prime / cadeau","La récompense offerte (produit, accessoire, bon…), généralement expédiée au participant."),
     ("Justificatif d'achat","Preuve d'achat demandée au participant (facture, ticket de caisse…)."),
     ("Adresse de livraison","Adresse à laquelle la prime (cadeau) sera expédiée."),
     ("Site gabarit","Version de test du site de l'opération, identique au site final, pour tout valider avant la mise en ligne."),
     ("Houston","Outil interne Promo.dev (la barre en haut du site de test) qui sert à simuler les états de l'offre et à prévisualiser les pages. Invisible pour les participants."),
     ("Dédoublonnage","Contrôle qui empêche les participations en double (ex. 1 prime par foyer / justificatif)."),
     ("Suivi de participation","Page/e-mail permettant au participant de suivre le statut de traitement de son dossier."),
     ("Bon pour mise en ligne","Validation finale du client attestant que l'opération peut être publiée.")]
    for term,desc in lex:
        tcell=ws.cell(r,2,term); tcell.font=F(11,True,"26324A"); tcell.fill=fill(GREY)
        tcell.alignment=Alignment(vertical="top",wrap_text=True,indent=1); bd(ws,tcell.coordinate)
        dcell=ws.cell(r,3,desc); dcell.font=F(11); dcell.alignment=Alignment(vertical="top",wrap_text=True,indent=1); bd(ws,dcell.coordinate)
        ws.row_dimensions[r].height=32; r+=1
    ws.page_setup.orientation="portrait"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ws.print_area=f"A1:D{r}"

    # ================= 2) GUIDE HOUSTON =================
    g=wb.create_sheet("Guide Houston"); g.sheet_view.showGridLines=False
    g.column_dimensions["A"].width=2.5; g.column_dimensions["B"].width=30
    g.column_dimensions["C"].width=60; g.column_dimensions["D"].width=60
    r=title_bar(g,2,4,2,"GUIDE DE LA BARRE D'ADMINISTRATION « HOUSTON »",
                "La barre en haut du site de test permet de simuler tous les statuts et pages de l'offre."); r+=1
    hh=g.cell(r,2,"⚠  La barre « Houston » est un outil de test interne"); hh.font=F(12,True,WHITE); hh.fill=fill(NAVY)
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); hh.alignment=leftc; g.row_dimensions[r].height=24; r+=1
    frm=("Elle apparaît uniquement sur ce site de test et sert à simuler les différents états de l'offre "
         "(en cours, en attente, terminé…) et à prévisualiser les pages. Elle ne sera JAMAIS visible par les participants sur le site final.")
    fc=g.cell(r,2,frm); fc.font=F(11); fc.alignment=Alignment(wrap_text=True,vertical="center",indent=1); fc.fill=fill("FBF3E6")
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    for cc in range(2,5): g.cell(r,cc).border=border
    g.row_dimensions[r].height=48; r+=2
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    for cc in range(2,5): g.cell(r,cc).fill=fill(TINT); g.cell(r,cc).border=border
    place_abs(g,"image9.png",2,r,720,90,3,90); r+=2
    qh=g.cell(r,2,"  PAR OÙ COMMENCER — testez en 6 étapes"); qh.font=F(12,True,WHITE); qh.fill=fill(ACCENT)
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); qh.alignment=leftc; g.row_dimensions[r].height=22; r+=1
    qs=["Ouvrez l'URL de test (voir l'onglet « Mode d'emploi »).",
        "Dans la barre Houston, réglez « État » sur « En cours ».",
        "Cliquez sur « Participer » et remplissez le formulaire (identité + adresse de livraison).",
        "Téléversez le justificatif d'achat, puis validez la participation.",
        "Vérifiez la page de confirmation et les e-mails reçus (confirmation, puis validation du dossier).",
        "Notez vos remarques dans « Checklist PRIME » (OK/KO) et détaillez tout souci dans « Anomalies »."]
    for i,s in enumerate(qs,1):
        zeb=ZEBRA if i%2==0 else WHITE
        eb=g.cell(r,2,f"Étape {i}"); eb.font=F(11,True,ACCENT); eb.alignment=Alignment(vertical="top",indent=1); eb.fill=fill(zeb); bd(g,eb.coordinate)
        sc=g.cell(r,3,s); sc.font=F(11); sc.alignment=wrap; sc.fill=fill(zeb)
        g.merge_cells(start_row=r,start_column=3,end_row=r,end_column=4); g.cell(r,4).fill=fill(zeb)
        bd(g,f"C{r}"); bd(g,f"D{r}")
        g.row_dimensions[r].height=42; r+=1
    r+=1
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    nc=g.cell(r,2,"ℹ Les aperçus ci-dessous sont des exemples génériques. La barre d'administration Houston est identique pour toutes les opérations ; seules les pages de l'offre changent selon la marque.")
    nc.font=F(10,True,"6B7280"); nc.fill=fill("F2F4F8"); nc.alignment=Alignment(wrap_text=True,vertical="center",indent=1)
    for cc in range(2,5): g.cell(r,cc).border=border
    g.row_dimensions[r].height=32; r+=2
    hr=r
    for i,txt in enumerate(("CONTRÔLE","À QUOI ÇA SERT","APERÇU")):
        c=g.cell(hr,2+i,txt); c.font=F(11,True,WHITE); c.fill=fill(BLUE); c.alignment=wrapc; bd(g,c.coordinate)
    g.row_dimensions[hr].height=24
    guide=[("CONNEXION","Se connecter sur l'URL de test de l'opération (voir onglet « Mode d'emploi »). Un justificatif d'achat de test est nécessaire pour aller au bout du parcours.",None,0),
     ("ÉTAT","Permet de tester les différents statuts de l'opération (menu déroulant).","image6.png",300),
     ("ÉTAT ▸ En cours","Le parcours est accessible ; le bouton « Participer » apparaît pour lancer et valider une participation.","image1.png",210),
     ("ÉTAT ▸ En attente","Affiche la page d'attente : l'offre n'a pas encore commencé.","neutral_attente.png",380),
     ("ÉTAT ▸ Terminé","Affiche la page de fin : l'offre est terminée.","neutral_termine.png",380),
     ("ÉTAT ▸ Offline","Affiche la page « hors ligne » quand l'offre est suspendue.","image12.png",380),
     ("ÉTAT ▸ Quota atteint","Affiche la page d'accueil quand le nombre max de participations est atteint (si compteur limitatif).","image13.png",380),
     ("VISUEL","Permet de visualiser directement la Home et — pour chaque prime — sa page de confirmation (le menu liste le nom de chaque cadeau).","visuel_prime.png",400),
     ("PARTICIPATION (statut)","PERD ou GAGNE. Sur une offre de prime, sans incidence sur le parcours ni sur la page de confirmation (ce n'est pas un jeu mais une prime).","image8.png",240),
     ("DÉSACTIVER CONTRÔLES","Désactive le contrôle d'unicité : permet de retester plusieurs fois avec le même e-mail / justificatif. Décoché = le dédoublonnage paramétré s'applique.","image3.png",260),
     ("SUIVI DE PARTICIPATION","Depuis le pied de page, permet de recevoir un e-mail avec le lien de suivi du dossier (statut de traitement).","neutral_suivi.png",360)]
    row=hr+1
    for idx,(label,desc,imgn,tw) in enumerate(guide):
        zeb=ZEBRA if idx%2==0 else WHITE
        lc=g.cell(row,2,label); lc.font=F(11,True,NAVY); lc.alignment=wrapv; lc.fill=fill(TINT); bd(g,lc.coordinate)
        dc=g.cell(row,3,desc); dc.font=F(12); dc.alignment=wrap; dc.fill=fill(zeb); bd(g,dc.coordinate)
        ac=g.cell(row,4); ac.fill=fill(zeb); bd(g,ac.coordinate)
        if imgn:
            place_abs(g,imgn,4,row,tw,230,1,60)
        else:
            g.row_dimensions[row].height=42
        row+=1
    g.page_setup.orientation="landscape"; g.page_setup.fitToWidth=1; g.page_setup.fitToHeight=0
    g.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); g.print_area=f"A1:D{row}"

    # ================= 3) CHECKLIST PRIME =================
    ck=wb.create_sheet("Checklist PRIME"); ck.sheet_view.showGridLines=False
    for col,w in {"A":2.5,"B":5.5,"C":23,"D":46,"E":46,"F":12,"G":12,"H":12,"I":34}.items(): ck.column_dimensions[col].width=w
    r=title_bar(ck,2,9,2,"CHECKLIST DE TEST — OFFRE DE PRIME (PRIME)",
                "Renseignez un statut par appareil (Desktop / Tablette / Mobile). Marquez « N.A. » les tests non applicables. Tout KO doit être détaillé dans « Anomalies »."); r+=1
    headers=["N°","Catégorie","Test à réaliser","Résultat attendu","Desktop","Tablette","Mobile","Commentaire / Anomalie"]
    hr=r
    for i,h in enumerate(headers):
        c=ck.cell(hr,2+i,h); c.font=F(11,True,WHITE); c.fill=fill(BLUE); c.alignment=wrapc; bd(ck,c.coordinate)
    ck.row_dimensions[hr].height=32
    CATCOL={"Accès & états":"E7EEF9","Home page":"E9F5EC","Formulaire":"FBF0E4","Adresse de livraison":"E4F3F5",
     "Justificatif d'achat":"F3EAF7","Preuve produit":"FDECEC","Validation":"EAF0FB","Consentements":"F0ECF9",
     "Anti-fraude":"ECEFF2","E-mails":"FFF7E0","Suivi":"E9F0FA"}
    tests=[("Accès & états","Accéder à l'URL de test","La page se charge, la barre Houston est visible en haut"),
     ("Accès & états","État « En attente »","La page d'attente s'affiche (offre pas encore commencée)"),
     ("Accès & états","État « En cours »","Le bouton « Participer » apparaît, le parcours est accessible"),
     ("Accès & états","État « Terminé »","La page de fin d'offre s'affiche"),
     ("Accès & états","État « Offline »","La page « opération suspendue » s'affiche"),
     ("Accès & états","État « Quota atteint » (si compteur)","La page « limite de participations atteinte » s'affiche"),
     ("Home page","Affichage Home desktop","Visuels, logos, textes et charte conformes à la maquette validée"),
     ("Home page","Affichage Home mobile (responsive)","Mise en page correcte sur smartphone, aucun élément coupé"),
     ("Home page","Modalités / mentions légales / confidentialité","Liens présents et documents accessibles"),
     ("Home page","Bouton « Participer »","Lance le parcours de participation"),
     ("Formulaire","Présence de tous les champs attendus","Civilité, nom, prénom, e-mail, confirmation e-mail, téléphone présents"),
     ("Formulaire","Champs obligatoires laissés vides","Message d'erreur clair, blocage de la validation"),
     ("Formulaire","Contrôles de format (e-mail, confirmation, tél.)","Une saisie invalide (ou e-mails différents) déclenche un message d'erreur"),
     ("Adresse de livraison","Saisie de l'adresse (autocomplétion)","La recherche d'adresse propose des suggestions et remplit les champs"),
     ("Adresse de livraison","« Vous ne trouvez pas votre adresse ? »","La saisie manuelle de l'adresse est possible"),
     ("Justificatif d'achat","Upload d'un justificatif valide (PDF/JPG/PNG)","Le fichier est accepté et visible dans le récapitulatif"),
     ("Justificatif d'achat","Upload d'un format / poids non autorisé","Message d'erreur, fichier refusé"),
     ("Justificatif d'achat","Validation sans justificatif","Blocage tant que le justificatif obligatoire n'est pas fourni"),
     ("Preuve produit","Numéro de série / IMEI / code produit (si demandé)","Le champ est présent et contrôlé ; une valeur invalide est refusée"),
     ("Validation","Récapitulatif avant envoi","Les données saisies (identité, adresse, justificatif) sont exactes et modifiables"),
     ("Validation","Valider la participation","Page de confirmation affichée (visuel + message de succès)"),
     ("Consentements","Case CGU / RGPD non cochée","Impossible de valider tant que le consentement obligatoire n'est pas donné"),
     ("Anti-fraude","Dédoublonnage (contrôles activés)","2e participation même e-mail / justificatif refusée avec message"),
     ("Anti-fraude","Désactiver contrôles","Permet de re-tester avec le même e-mail / justificatif"),
     ("E-mails","E-mail de confirmation de participation","E-mail reçu ; expéditeur, objet, visuels et liens corrects"),
     ("E-mails","E-mail de validation / refus du dossier","Le changement de statut déclenche le bon e-mail (dossier validé ou justificatif non conforme)"),
     ("E-mails","E-mail d'expédition de la prime","Un e-mail informe de l'envoi du cadeau (le moment venu)"),
     ("E-mails","Rendu e-mail sur mobile","E-mails lisibles et bien affichés sur smartphone"),
     ("Suivi","Suivi de participation (pied de page)","Après saisie de l'e-mail, réception du lien de suivi et bon statut affiché")]
    row=hr+1
    ex=("Ex","Justificatif d'achat","Téléverser une facture PDF valide","Fichier accepté, visible dans le récap","OK","OK","KO","KO mobile : bouton d'upload coupé")
    for i,v in enumerate(ex):
        c=ck.cell(row,2+i,v); c.border=border; c.alignment=(wrapc if i in(0,4,5,6) else wrap); c.font=F(10,it=True,color="9AA0AA"); c.fill=fill("FBFBFD")
    ck.row_dimensions[row].height=28; row+=1
    start=row; n=1
    for cat,test,exp in tests:
        zeb=ZEBRA if n%2==0 else WHITE
        a=ck.cell(row,2,n); a.alignment=center; a.font=F(11,True,"6B7280"); a.fill=fill(zeb); bd(ck,a.coordinate)
        cc=ck.cell(row,3,cat); cc.font=F(11,True,"2C3E63"); cc.alignment=wrapv; cc.fill=fill(CATCOL.get(cat,"EDEFF3")); bd(ck,cc.coordinate)
        t=ck.cell(row,4,test); t.font=F(11); t.alignment=wrap; t.fill=fill(zeb); bd(ck,t.coordinate)
        e=ck.cell(row,5,exp); e.font=F(11,color="4B5563"); e.alignment=wrap; e.fill=fill(zeb); bd(ck,e.coordinate)
        for col in (6,7,8):
            x=ck.cell(row,col); x.fill=fill(zeb); x.alignment=center; bd(ck,x.coordinate); x.font=F(11)
        cm=ck.cell(row,9); cm.fill=fill(zeb); cm.alignment=wrap; bd(ck,cm.coordinate); cm.font=F(11)
        ck.row_dimensions[row].height=34; row+=1; n+=1
    end=row-1
    dv=DataValidation(type="list",formula1='"OK,KO,N.A.,À tester"',allow_blank=True); dv.add(f"F{start}:H{end}"); ck.add_data_validation(dv)
    for val,bg,tx in [("OK",GREENF,GREENT),("KO",REDF,REDT),("N.A.",NEUTF,"555555"),("À tester",YELLOW,"B45309")]:
        ck.conditional_formatting.add(f"F{start}:H{end}",CellIsRule(operator="equal",formula=[f'"{val}"'],fill=fill(bg),font=F(11,True,tx)))
    sr=end+2
    ck.merge_cells(start_row=sr,start_column=3,end_row=sr,end_column=9)
    hh=ck.cell(sr,3,"  RÉCAPITULATIF (nombre de tests par statut et par appareil)"); hh.font=F(12,True,WHITE); hh.fill=fill(NAVY); hh.alignment=leftc; ck.row_dimensions[sr].height=24
    ck.merge_cells(start_row=sr+1,start_column=3,end_row=sr+1,end_column=5)
    for cidx,dev in [(6,"Desktop"),(7,"Tablette"),(8,"Mobile")]:
        dc=ck.cell(sr+1,cidx,dev); dc.font=F(11,True,NAVY); dc.fill=fill(TINT); dc.alignment=center; bd(ck,dc.coordinate)
    rr=sr+2
    for lab,kind,bg,tx in [("Tests OK","OK",GREENF,GREENT),("Tests KO","KO",REDF,REDT),("Restant à tester","REST",YELLOW,"B45309")]:
        ck.merge_cells(start_row=rr,start_column=3,end_row=rr,end_column=5)
        l=ck.cell(rr,3,lab); l.font=F(11,True,tx); l.fill=fill(bg); l.alignment=leftc; bd(ck,l.coordinate)
        for cidx,cl in [(6,"F"),(7,"G"),(8,"H")]:
            if kind=="REST":
                form=f'=COUNTA(D{start}:D{end})-COUNTIF({cl}{start}:{cl}{end},"OK")-COUNTIF({cl}{start}:{cl}{end},"KO")-COUNTIF({cl}{start}:{cl}{end},"N.A.")'
            else:
                form=f'=COUNTIF({cl}{start}:{cl}{end},"{kind}")'
            v=ck.cell(rr,cidx,form); v.font=F(12,True,tx); v.fill=fill(bg); v.alignment=center; bd(ck,v.coordinate)
        ck.row_dimensions[rr].height=22; rr+=1
    ck.freeze_panes=f"B{hr+1}"
    ck.page_setup.orientation="landscape"; ck.page_setup.fitToWidth=1; ck.page_setup.fitToHeight=0
    ck.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ck.print_area=f"A1:I{rr}"

    # ================= 4) ANOMALIES =================
    an=wb.create_sheet("Anomalies"); an.sheet_view.showGridLines=False
    for c,w in {"A":2.5,"B":5.5,"C":14,"D":24,"E":42,"F":36,"G":14,"H":14,"I":24}.items(): an.column_dimensions[c].width=w
    r=title_bar(an,2,9,2,"JOURNAL DES ANOMALIES",
                "Une ligne par bug. Reliez chaque anomalie au N° de test concerné et joignez une capture si possible."); r+=1
    hdr=["N°","N° test lié","Page / Écran","Description du problème","Étapes pour reproduire","Gravité","Statut","Capture (nom fichier)"]
    hr=r
    for i,h in enumerate(hdr):
        c=an.cell(hr,2+i,h); c.font=F(11,True,WHITE); c.fill=fill(BLUE); c.alignment=wrapc; bd(an,c.coordinate)
    an.row_dimensions[hr].height=32
    exa=("Ex","16","Justificatif","Le bouton d'upload est coupé sur mobile","1. Aller à l'étape justificatif  2. Ouvrir sur smartphone","Majeure","Ouvert","capture_01.png")
    row=hr+1
    for i,v in enumerate(exa):
        c=an.cell(row,2+i,v); c.border=border; c.alignment=(wrapc if i in(0,1,5,6) else wrap); c.font=F(10,it=True,color="9AA0AA"); c.fill=fill("FBFBFD")
    an.row_dimensions[row].height=30; row+=1
    first=row
    for k in range(20):
        zeb=ZEBRA if k%2==0 else WHITE
        for col in range(2,10):
            c=an.cell(row,col); c.border=border; c.font=F(11); c.fill=fill(zeb); c.alignment=(wrapc if col in(2,3,7,8) else wrap)
        an.row_dimensions[row].height=28; row+=1
    last=row-1
    dvg=DataValidation(type="list",formula1='"Bloquante,Majeure,Mineure,Cosmétique"',allow_blank=True); dvg.add(f"G{first}:G{last}"); an.add_data_validation(dvg)
    dvs=DataValidation(type="list",formula1='"Ouvert,Corrigé,Vérifié,Rejeté"',allow_blank=True); dvs.add(f"H{first}:H{last}"); an.add_data_validation(dvs)
    an.conditional_formatting.add(f"G{first}:G{last}",CellIsRule(operator="equal",formula=['"Bloquante"'],fill=fill(REDF),font=F(11,True,REDT)))
    an.conditional_formatting.add(f"G{first}:G{last}",CellIsRule(operator="equal",formula=['"Majeure"'],fill=fill("FCE5CD"),font=F(11,True,"B45309")))
    an.conditional_formatting.add(f"H{first}:H{last}",CellIsRule(operator="equal",formula=['"Corrigé"'],fill=fill(GREENF),font=F(11,True,GREENT)))
    an.conditional_formatting.add(f"H{first}:H{last}",CellIsRule(operator="equal",formula=['"Vérifié"'],fill=fill(GREENF),font=F(11,True,GREENT)))
    an.freeze_panes=f"B{hr+1}"
    an.page_setup.orientation="landscape"; an.page_setup.fitToWidth=1; an.page_setup.fitToHeight=0
    an.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); an.print_area=f"A1:I{last}"

    # ================= 5) VALIDATION =================
    va=wb.create_sheet("Validation"); va.sheet_view.showGridLines=False
    va.column_dimensions["A"].width=2.5; va.column_dimensions["B"].width=44; va.column_dimensions["C"].width=54; va.column_dimensions["D"].width=3
    r=title_bar(va,2,3,2,"VALIDATION — BON POUR MISE EN LIGNE",
                "À compléter une fois la checklist terminée. Cette validation vaut accord pour la mise en ligne de l'opération."); r+=1
    valrow=None
    for label,val in [("Opération validée","__DROPDOWN__"),
     ("Réserves / points en suspens",""),("Nombre d'anomalies bloquantes restantes",""),
     ("Nom du valideur (client)",""),("Fonction",""),("Date de validation",""),
     ("Signature (nom + « bon pour accord »)",""),("Chef de projet Promo.dev","")]:
        l=va.cell(r,2,label); l.font=F(11,True,"26324A"); l.fill=fill(GREY); l.alignment=leftc; bd(va,l.coordinate)
        c=va.cell(r,3); c.fill=fill(YELLOW); c.alignment=Alignment(vertical="center",wrap_text=True,indent=1)
        yb=Side(style="thin",color=YELLOWB); c.border=Border(left=yb,right=yb,top=yb,bottom=yb)
        if val=="__DROPDOWN__":
            c.value="— cliquez pour choisir —"; c.font=F(11,it=True,color="9AA0AA"); valrow=r
        else:
            c.value=val; c.font=F(11)
        va.row_dimensions[r].height=38 if val=="" else 28; r+=1
    dvv=DataValidation(type="list",formula1='"Oui sans réserve,Oui avec réserves,Non"',allow_blank=True)
    dvv.add(f"C{valrow}"); va.add_data_validation(dvv)
    va.conditional_formatting.add(f"C{valrow}",CellIsRule(operator="equal",formula=['"Oui sans réserve"'],fill=fill(GREENF),font=F(11,True,GREENT)))
    va.conditional_formatting.add(f"C{valrow}",CellIsRule(operator="equal",formula=['"Oui avec réserves"'],fill=fill(YELLOW),font=F(11,True,"B45309")))
    va.conditional_formatting.add(f"C{valrow}",CellIsRule(operator="equal",formula=['"Non"'],fill=fill(REDF),font=F(11,True,REDT)))
    va.page_setup.orientation="portrait"; va.page_setup.fitToWidth=1; va.page_setup.fitToHeight=0
    va.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); va.print_area=f"A1:D{r}"


def build_formulaire(wb):


    # ================= 1) MODE D'EMPLOI =================
    ws=wb.active; ws.title="Mode d'emploi"; ws.sheet_view.showGridLines=False
    ws.column_dimensions["A"].width=2.5; ws.column_dimensions["B"].width=40
    ws.column_dimensions["C"].width=76; ws.column_dimensions["D"].width=3
    r=title_bar(ws,2,3,2,"PROCÉDURE DE TEST — FORMULAIRE (COLLECTE / OPT-IN)",
                "Gabarit générique · à dupliquer et compléter pour chaque opération"); r+=1
    hc=ws.cell(r,2,"  À LIRE EN PREMIER"); hc.font=F(12,True,WHITE); hc.fill=fill(ACCENT)
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3); hc.alignment=leftc; ws.row_dimensions[r].height=22; r+=1
    intro=("Vous allez tester le site de votre formulaire (inscription, collecte de données, opt-in newsletter…) avant sa mise en ligne. "
           "Ce site « gabarit » est une copie de test, identique à ce que verront les participants. "
           "Objectif : vérifier le parcours, les champs du formulaire, les consentements (RGPD) et l'e-mail de confirmation, puis valider avant le lancement.\n\n"
           "Comptez environ 20 minutes. Avant de commencer, munissez-vous d'une boîte mail que vous pouvez consulter (pensez à vérifier les spams). "
           "Testez de préférence sur ordinateur puis sur mobile. En cas de blocage, contactez votre chef de projet Promo.dev (coordonnées ci-dessous).")
    ic=ws.cell(r,2,intro); ic.font=F(11); ic.alignment=Alignment(wrap_text=True,vertical="top",indent=1); ic.fill=fill("FBF3E6")
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    for cc in (2,3): ws.cell(r,cc).border=border
    ws.row_dimensions[r].height=140; r+=2
    section(ws,2,r,"1 · Informations de l'opération"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    tag=ws.cell(r,2,"▸ À compléter par le chef de projet Promo.dev avant l'envoi au client")
    tag.font=F(10,True,ACCENT); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    tag.alignment=Alignment(indent=1,vertical="center"); ws.row_dimensions[r].height=20; r+=1
    info=[("Client / Marque","ex. Marque"),("Nom de l'opération","ex. Inscription Newsletter"),
     ("N° opération (idgame)","ex. 999"),("URL de test","ex. https://test.promo.dev/xxxxxxxx"),
     ("Objet du formulaire","ex. inscription newsletter / collecte de données"),
     ("Champs collectés","ex. civilité, nom, prénom, e-mail, code postal"),
     ("Pièce jointe demandée ?","ex. non  /  oui : justificatif à téléverser"),
     ("Double opt-in (confirmation e-mail) ?","ex. oui : lien à cliquer  /  non"),
     ("E-mail de test","ex. prenom.nom+test@domaine.fr"),
     ("Dédoublonnage paramétré","ex. 1 inscription / e-mail"),
     ("Dates de l'opération","ex. 01/09/2026 → 31/12/2026"),
     ("Testeur(s) côté client","ex. Prénom Nom"),("Contact Promo.dev (chef de projet)","ex. prenom@promo.dev"),
     ("Date limite de retour des tests","ex. JJ/MM/AAAA")]
    for label,exv in info:
        lc=ws.cell(r,2,label); lc.font=F(11,True,"26324A"); lc.fill=fill(GREY); lc.alignment=leftc; bd(ws,lc.coordinate)
        val=INFO.get(label,exv); is_real=label in INFO
        vc=ws.cell(r,3,val); vc.font=F(11,False,"26324A" if is_real else "9AA0AA"); vc.fill=fill(YELLOW); vc.alignment=leftc
        yb=Side(style="thin",color=YELLOWB); vc.border=Border(left=yb,right=yb,top=yb,bottom=yb)
        ws.row_dimensions[r].height=26; r+=1
    r+=1
    section(ws,2,r,"2 · Comment utiliser ce classeur"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    steps=["Vérifiez les informations de l'opération ci-dessus (pré-remplies par Promo.dev) : URL, objet, champs, dates.",
     "Onglet « Guide Houston » : familiarisez-vous avec la barre d'administration et les états de l'opération.",
     "Onglet « Checklist FORMULAIRE » : déroulez chaque test, indiquez un statut par appareil (OK / KO / N.A.) et un commentaire.",
     "Onglet « Anomalies » : décrivez précisément chaque bug (page, action, capture).",
     "Onglet « Validation » : donnez votre bon pour mise en ligne une fois tous les tests OK.",
     "Renvoyez le fichier complété au chef de projet avant la date limite."]
    for i,s in enumerate(steps,1):
        b=ws.cell(r,2,f"Étape {i}"); b.font=F(11,True,ACCENT); b.alignment=Alignment(vertical="center",indent=1)
        b.fill=fill(ZEBRA if i%2==0 else WHITE)
        c=ws.cell(r,3,s); c.font=F(11); c.alignment=wrapv; c.fill=fill(ZEBRA if i%2==0 else WHITE)
        ws.row_dimensions[r].height=32; r+=1
    r+=1
    section(ws,2,r,"3 · Légende des statuts"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    for code,desc,bg,tx in [("OK","Conforme, rien à signaler",GREENF,GREENT),
     ("KO","Anomalie détectée → à détailler dans l'onglet Anomalies",REDF,REDT),
     ("N.A.","Non applicable à cette opération",NEUTF,"555555"),("À tester","Test pas encore réalisé",YELLOW,"B45309")]:
        p=ws.cell(r,2,code); p.font=F(11,True,tx); p.fill=fill(bg); p.alignment=center; bd(ws,p.coordinate)
        d=ws.cell(r,3,desc); d.font=F(11); d.alignment=leftc; ws.row_dimensions[r].height=22; r+=1
    r+=1
    section(ws,2,r,"4 · Lexique"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    lex=[("Formulaire (collecte / opt-in)","Opération de recueil de données (inscription, newsletter, sondage…), sans achat ni lot à gagner."),
     ("Opt-in","Consentement explicite du participant pour être recontacté (ex. inscription newsletter)."),
     ("Double opt-in","Confirmation de l'inscription via un lien envoyé par e-mail au participant."),
     ("Consentement RGPD","Accord obligatoire au traitement des données personnelles, avec mention d'information."),
     ("Site gabarit","Version de test du site de l'opération, identique au site final, pour tout valider avant la mise en ligne."),
     ("Houston","Outil interne Promo.dev (la barre en haut du site de test) qui sert à simuler les états et prévisualiser les pages. Invisible pour les participants."),
     ("Dédoublonnage","Contrôle qui empêche les inscriptions en double (ex. 1 par e-mail)."),
     ("Suivi de participation","Page/e-mail permettant au participant de retrouver le statut de sa demande."),
     ("Bon pour mise en ligne","Validation finale du client attestant que l'opération peut être publiée.")]
    for term,desc in lex:
        tcell=ws.cell(r,2,term); tcell.font=F(11,True,"26324A"); tcell.fill=fill(GREY)
        tcell.alignment=Alignment(vertical="top",wrap_text=True,indent=1); bd(ws,tcell.coordinate)
        dcell=ws.cell(r,3,desc); dcell.font=F(11); dcell.alignment=Alignment(vertical="top",wrap_text=True,indent=1); bd(ws,dcell.coordinate)
        ws.row_dimensions[r].height=32; r+=1
    ws.page_setup.orientation="portrait"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ws.print_area=f"A1:D{r}"

    # ================= 2) GUIDE HOUSTON =================
    g=wb.create_sheet("Guide Houston"); g.sheet_view.showGridLines=False
    g.column_dimensions["A"].width=2.5; g.column_dimensions["B"].width=30
    g.column_dimensions["C"].width=60; g.column_dimensions["D"].width=60
    r=title_bar(g,2,4,2,"GUIDE DE LA BARRE D'ADMINISTRATION « HOUSTON »",
                "La barre en haut du site de test permet de simuler tous les statuts et pages de l'opération."); r+=1
    hh=g.cell(r,2,"⚠  La barre « Houston » est un outil de test interne"); hh.font=F(12,True,WHITE); hh.fill=fill(NAVY)
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); hh.alignment=leftc; g.row_dimensions[r].height=24; r+=1
    frm=("Elle apparaît uniquement sur ce site de test et sert à simuler les différents états de l'opération "
         "(en cours, en attente, terminé…) et à prévisualiser les pages. Elle ne sera JAMAIS visible par les participants sur le site final.")
    fc=g.cell(r,2,frm); fc.font=F(11); fc.alignment=Alignment(wrap_text=True,vertical="center",indent=1); fc.fill=fill("FBF3E6")
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    for cc in range(2,5): g.cell(r,cc).border=border
    g.row_dimensions[r].height=48; r+=2
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    for cc in range(2,5): g.cell(r,cc).fill=fill(TINT); g.cell(r,cc).border=border
    place_abs(g,"image9.png",2,r,720,90,3,90); r+=2
    qh=g.cell(r,2,"  PAR OÙ COMMENCER — testez en 6 étapes"); qh.font=F(12,True,WHITE); qh.fill=fill(ACCENT)
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); qh.alignment=leftc; g.row_dimensions[r].height=22; r+=1
    qs=["Ouvrez l'URL de test (voir l'onglet « Mode d'emploi »).",
        "Dans la barre Houston, réglez « État » sur « En cours ».",
        "Cliquez sur « Participer / S'inscrire » et remplissez le formulaire.",
        "Cochez les consentements (RGPD) et validez.",
        "Vérifiez la page de confirmation et l'e-mail de confirmation (double opt-in : cliquez le lien reçu, si prévu).",
        "Notez vos remarques dans « Checklist FORMULAIRE » (OK/KO) et détaillez tout souci dans « Anomalies »."]
    for i,s in enumerate(qs,1):
        zeb=ZEBRA if i%2==0 else WHITE
        eb=g.cell(r,2,f"Étape {i}"); eb.font=F(11,True,ACCENT); eb.alignment=Alignment(vertical="top",indent=1); eb.fill=fill(zeb); bd(g,eb.coordinate)
        sc=g.cell(r,3,s); sc.font=F(11); sc.alignment=wrap; sc.fill=fill(zeb)
        g.merge_cells(start_row=r,start_column=3,end_row=r,end_column=4); g.cell(r,4).fill=fill(zeb)
        bd(g,f"C{r}"); bd(g,f"D{r}")
        g.row_dimensions[r].height=42; r+=1
    r+=1
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    nc=g.cell(r,2,"ℹ Les aperçus ci-dessous sont des exemples génériques. La barre d'administration Houston est identique pour toutes les opérations ; seules les pages changent selon la marque.")
    nc.font=F(10,True,"6B7280"); nc.fill=fill("F2F4F8"); nc.alignment=Alignment(wrap_text=True,vertical="center",indent=1)
    for cc in range(2,5): g.cell(r,cc).border=border
    g.row_dimensions[r].height=32; r+=2
    hr=r
    for i,txt in enumerate(("CONTRÔLE","À QUOI ÇA SERT","APERÇU")):
        c=g.cell(hr,2+i,txt); c.font=F(11,True,WHITE); c.fill=fill(BLUE); c.alignment=wrapc; bd(g,c.coordinate)
    g.row_dimensions[hr].height=24
    guide=[("CONNEXION","Se connecter sur l'URL de test de l'opération (voir onglet « Mode d'emploi »).",None,0),
     ("ÉTAT","Permet de tester les différents statuts de l'opération (menu déroulant).","image6.png",300),
     ("ÉTAT ▸ En cours","Le parcours est accessible ; le bouton « Participer / S'inscrire » apparaît pour lancer et valider une participation.","image1.png",210),
     ("ÉTAT ▸ En attente","Affiche la page d'attente : l'opération n'a pas encore commencé.","neutral_attente_op.png",380),
     ("ÉTAT ▸ Terminé","Affiche la page de fin : l'opération est terminée.","neutral_termine_op.png",380),
     ("ÉTAT ▸ Offline","Affiche la page « hors ligne » quand l'opération est suspendue.","image12.png",380),
     ("ÉTAT ▸ Quota atteint","Affiche la page d'accueil quand le nombre max de participations est atteint (si compteur limitatif).","image13.png",380),
     ("VISUEL","Permet de visualiser la Home et la page de confirmation d'inscription.","visuel_form.png",400),
     ("PARTICIPATION (statut)","PERD ou GAGNE. Sur un formulaire, sans incidence sur le parcours ni la page de confirmation (ce n'est pas un jeu).","image8.png",240),
     ("DÉSACTIVER CONTRÔLES","Désactive le contrôle d'unicité : permet de retester plusieurs fois avec le même e-mail. Décoché = le dédoublonnage paramétré s'applique (ex. 1 inscription / e-mail).","image3.png",260),
     ("SUIVI DE PARTICIPATION","Depuis le pied de page, permet de recevoir un e-mail avec le lien de suivi de la demande.","neutral_suivi.png",360)]
    row=hr+1
    for idx,(label,desc,imgn,tw) in enumerate(guide):
        zeb=ZEBRA if idx%2==0 else WHITE
        lc=g.cell(row,2,label); lc.font=F(11,True,NAVY); lc.alignment=wrapv; lc.fill=fill(TINT); bd(g,lc.coordinate)
        dc=g.cell(row,3,desc); dc.font=F(12); dc.alignment=wrap; dc.fill=fill(zeb); bd(g,dc.coordinate)
        ac=g.cell(row,4); ac.fill=fill(zeb); bd(g,ac.coordinate)
        if imgn:
            place_abs(g,imgn,4,row,tw,230,1,60)
        else:
            g.row_dimensions[row].height=42
        row+=1
    g.page_setup.orientation="landscape"; g.page_setup.fitToWidth=1; g.page_setup.fitToHeight=0
    g.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); g.print_area=f"A1:D{row}"

    # ================= 3) CHECKLIST FORMULAIRE =================
    ck=wb.create_sheet("Checklist FORMULAIRE"); ck.sheet_view.showGridLines=False
    for col,w in {"A":2.5,"B":5.5,"C":23,"D":46,"E":46,"F":12,"G":12,"H":12,"I":34}.items(): ck.column_dimensions[col].width=w
    r=title_bar(ck,2,9,2,"CHECKLIST DE TEST — FORMULAIRE (COLLECTE / OPT-IN)",
                "Renseignez un statut par appareil (Desktop / Tablette / Mobile). Marquez « N.A. » les tests non applicables (ex. pièce jointe, double opt-in). Tout KO doit être détaillé dans « Anomalies »."); r+=1
    headers=["N°","Catégorie","Test à réaliser","Résultat attendu","Desktop","Tablette","Mobile","Commentaire / Anomalie"]
    hr=r
    for i,h in enumerate(headers):
        c=ck.cell(hr,2+i,h); c.font=F(11,True,WHITE); c.fill=fill(BLUE); c.alignment=wrapc; bd(ck,c.coordinate)
    ck.row_dimensions[hr].height=32
    CATCOL={"Accès & états":"E7EEF9","Home page":"E9F5EC","Formulaire":"FBF0E4","Pièce jointe":"F3EAF7",
     "Consentements RGPD":"FDECEC","Validation":"EAF0FB","E-mails":"FFF7E0","Anti-fraude":"ECEFF2","Suivi":"E9F0FA"}
    tests=[("Accès & états","Accéder à l'URL de test","La page se charge, la barre Houston est visible en haut"),
     ("Accès & états","État « En attente »","La page d'attente s'affiche (opération pas encore commencée)"),
     ("Accès & états","État « En cours »","Le bouton « Participer / S'inscrire » apparaît, le parcours est accessible"),
     ("Accès & états","État « Terminé »","La page de fin d'opération s'affiche"),
     ("Accès & états","État « Offline »","La page « opération suspendue » s'affiche"),
     ("Accès & états","État « Quota atteint » (si compteur)","La page « limite atteinte » s'affiche"),
     ("Home page","Affichage Home desktop","Visuels, logos, textes et charte conformes à la maquette validée"),
     ("Home page","Affichage Home mobile (responsive)","Mise en page correcte sur smartphone, aucun élément coupé"),
     ("Home page","Mentions légales / confidentialité","Liens présents et documents accessibles"),
     ("Home page","Bouton « Participer / S'inscrire »","Lance le parcours de participation"),
     ("Formulaire","Présence de tous les champs attendus","Tous les champs du cahier des charges sont présents"),
     ("Formulaire","Champs obligatoires laissés vides","Message d'erreur clair, blocage de la validation"),
     ("Formulaire","Contrôles de format (e-mail, code postal, tél.)","Une saisie invalide déclenche un message d'erreur"),
     ("Pièce jointe","Upload d'un document valide (si demandé)","Le fichier est accepté et visible dans le récapitulatif"),
     ("Pièce jointe","Upload d'un format / poids non autorisé (si demandé)","Message d'erreur, fichier refusé"),
     ("Consentements RGPD","Case de consentement obligatoire non cochée","Impossible de valider tant que le consentement obligatoire n'est pas donné"),
     ("Consentements RGPD","Opt-in marketing facultatif","La case marketing reste facultative (décochée par défaut) et n'empêche pas de valider"),
     ("Consentements RGPD","Mention d'information RGPD / notice données","La mention et le lien vers la notice de données sont présents"),
     ("Validation","Récapitulatif avant envoi (si présent)","Les données saisies sont exactes et modifiables"),
     ("Validation","Valider l'inscription","Page de confirmation affichée (message de succès)"),
     ("E-mails","E-mail de confirmation d'inscription","E-mail reçu ; expéditeur, objet, visuels et liens corrects"),
     ("E-mails","Double opt-in (si prévu)","Le lien de confirmation reçu par e-mail active bien l'inscription"),
     ("E-mails","Rendu e-mail sur mobile","E-mail lisible et bien affiché sur smartphone"),
     ("Anti-fraude","Dédoublonnage (contrôles activés)","2e inscription avec le même e-mail refusée avec message"),
     ("Anti-fraude","Désactiver contrôles","Permet de re-tester avec le même e-mail"),
     ("Suivi","Suivi de participation (pied de page)","Après saisie de l'e-mail, réception du lien de suivi et bon statut affiché")]
    row=hr+1
    ex=("Ex","Consentements RGPD","Valider sans cocher le consentement obligatoire","Message d'erreur, validation bloquée","OK","OK","KO","KO mobile : case masquée")
    for i,v in enumerate(ex):
        c=ck.cell(row,2+i,v); c.border=border; c.alignment=(wrapc if i in(0,4,5,6) else wrap); c.font=F(10,it=True,color="9AA0AA"); c.fill=fill("FBFBFD")
    ck.row_dimensions[row].height=28; row+=1
    start=row; n=1
    for cat,test,exp in tests:
        zeb=ZEBRA if n%2==0 else WHITE
        a=ck.cell(row,2,n); a.alignment=center; a.font=F(11,True,"6B7280"); a.fill=fill(zeb); bd(ck,a.coordinate)
        cc=ck.cell(row,3,cat); cc.font=F(11,True,"2C3E63"); cc.alignment=wrapv; cc.fill=fill(CATCOL.get(cat,"EDEFF3")); bd(ck,cc.coordinate)
        t=ck.cell(row,4,test); t.font=F(11); t.alignment=wrap; t.fill=fill(zeb); bd(ck,t.coordinate)
        e=ck.cell(row,5,exp); e.font=F(11,color="4B5563"); e.alignment=wrap; e.fill=fill(zeb); bd(ck,e.coordinate)
        for col in (6,7,8):
            x=ck.cell(row,col); x.fill=fill(zeb); x.alignment=center; bd(ck,x.coordinate); x.font=F(11)
        cm=ck.cell(row,9); cm.fill=fill(zeb); cm.alignment=wrap; bd(ck,cm.coordinate); cm.font=F(11)
        ck.row_dimensions[row].height=34; row+=1; n+=1
    end=row-1
    dv=DataValidation(type="list",formula1='"OK,KO,N.A.,À tester"',allow_blank=True); dv.add(f"F{start}:H{end}"); ck.add_data_validation(dv)
    for val,bg,tx in [("OK",GREENF,GREENT),("KO",REDF,REDT),("N.A.",NEUTF,"555555"),("À tester",YELLOW,"B45309")]:
        ck.conditional_formatting.add(f"F{start}:H{end}",CellIsRule(operator="equal",formula=[f'"{val}"'],fill=fill(bg),font=F(11,True,tx)))
    sr=end+2
    ck.merge_cells(start_row=sr,start_column=3,end_row=sr,end_column=9)
    hh=ck.cell(sr,3,"  RÉCAPITULATIF (nombre de tests par statut et par appareil)"); hh.font=F(12,True,WHITE); hh.fill=fill(NAVY); hh.alignment=leftc; ck.row_dimensions[sr].height=24
    ck.merge_cells(start_row=sr+1,start_column=3,end_row=sr+1,end_column=5)
    for cidx,dev in [(6,"Desktop"),(7,"Tablette"),(8,"Mobile")]:
        dc=ck.cell(sr+1,cidx,dev); dc.font=F(11,True,NAVY); dc.fill=fill(TINT); dc.alignment=center; bd(ck,dc.coordinate)
    rr=sr+2
    for lab,kind,bg,tx in [("Tests OK","OK",GREENF,GREENT),("Tests KO","KO",REDF,REDT),("Restant à tester","REST",YELLOW,"B45309")]:
        ck.merge_cells(start_row=rr,start_column=3,end_row=rr,end_column=5)
        l=ck.cell(rr,3,lab); l.font=F(11,True,tx); l.fill=fill(bg); l.alignment=leftc; bd(ck,l.coordinate)
        for cidx,cl in [(6,"F"),(7,"G"),(8,"H")]:
            if kind=="REST":
                form=f'=COUNTA(D{start}:D{end})-COUNTIF({cl}{start}:{cl}{end},"OK")-COUNTIF({cl}{start}:{cl}{end},"KO")-COUNTIF({cl}{start}:{cl}{end},"N.A.")'
            else:
                form=f'=COUNTIF({cl}{start}:{cl}{end},"{kind}")'
            v=ck.cell(rr,cidx,form); v.font=F(12,True,tx); v.fill=fill(bg); v.alignment=center; bd(ck,v.coordinate)
        ck.row_dimensions[rr].height=22; rr+=1
    ck.freeze_panes=f"B{hr+1}"
    ck.page_setup.orientation="landscape"; ck.page_setup.fitToWidth=1; ck.page_setup.fitToHeight=0
    ck.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ck.print_area=f"A1:I{rr}"

    # ================= 4) ANOMALIES =================
    an=wb.create_sheet("Anomalies"); an.sheet_view.showGridLines=False
    for c,w in {"A":2.5,"B":5.5,"C":14,"D":24,"E":42,"F":36,"G":14,"H":14,"I":24}.items(): an.column_dimensions[c].width=w
    r=title_bar(an,2,9,2,"JOURNAL DES ANOMALIES",
                "Une ligne par bug. Reliez chaque anomalie au N° de test concerné et joignez une capture si possible."); r+=1
    hdr=["N°","N° test lié","Page / Écran","Description du problème","Étapes pour reproduire","Gravité","Statut","Capture (nom fichier)"]
    hr=r
    for i,h in enumerate(hdr):
        c=an.cell(hr,2+i,h); c.font=F(11,True,WHITE); c.fill=fill(BLUE); c.alignment=wrapc; bd(an,c.coordinate)
    an.row_dimensions[hr].height=32
    exa=("Ex","16","Formulaire","La case de consentement RGPD est déjà cochée par défaut","1. Ouvrir le formulaire  2. Observer la case","Majeure","Ouvert","capture_01.png")
    row=hr+1
    for i,v in enumerate(exa):
        c=an.cell(row,2+i,v); c.border=border; c.alignment=(wrapc if i in(0,1,5,6) else wrap); c.font=F(10,it=True,color="9AA0AA"); c.fill=fill("FBFBFD")
    an.row_dimensions[row].height=30; row+=1
    first=row
    for k in range(20):
        zeb=ZEBRA if k%2==0 else WHITE
        for col in range(2,10):
            c=an.cell(row,col); c.border=border; c.font=F(11); c.fill=fill(zeb); c.alignment=(wrapc if col in(2,3,7,8) else wrap)
        an.row_dimensions[row].height=28; row+=1
    last=row-1
    dvg=DataValidation(type="list",formula1='"Bloquante,Majeure,Mineure,Cosmétique"',allow_blank=True); dvg.add(f"G{first}:G{last}"); an.add_data_validation(dvg)
    dvs=DataValidation(type="list",formula1='"Ouvert,Corrigé,Vérifié,Rejeté"',allow_blank=True); dvs.add(f"H{first}:H{last}"); an.add_data_validation(dvs)
    an.conditional_formatting.add(f"G{first}:G{last}",CellIsRule(operator="equal",formula=['"Bloquante"'],fill=fill(REDF),font=F(11,True,REDT)))
    an.conditional_formatting.add(f"G{first}:G{last}",CellIsRule(operator="equal",formula=['"Majeure"'],fill=fill("FCE5CD"),font=F(11,True,"B45309")))
    an.conditional_formatting.add(f"H{first}:H{last}",CellIsRule(operator="equal",formula=['"Corrigé"'],fill=fill(GREENF),font=F(11,True,GREENT)))
    an.conditional_formatting.add(f"H{first}:H{last}",CellIsRule(operator="equal",formula=['"Vérifié"'],fill=fill(GREENF),font=F(11,True,GREENT)))
    an.freeze_panes=f"B{hr+1}"
    an.page_setup.orientation="landscape"; an.page_setup.fitToWidth=1; an.page_setup.fitToHeight=0
    an.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); an.print_area=f"A1:I{last}"

    # ================= 5) VALIDATION =================
    va=wb.create_sheet("Validation"); va.sheet_view.showGridLines=False
    va.column_dimensions["A"].width=2.5; va.column_dimensions["B"].width=44; va.column_dimensions["C"].width=54; va.column_dimensions["D"].width=3
    r=title_bar(va,2,3,2,"VALIDATION — BON POUR MISE EN LIGNE",
                "À compléter une fois la checklist terminée. Cette validation vaut accord pour la mise en ligne de l'opération."); r+=1
    valrow=None
    for label,val in [("Opération validée","__DROPDOWN__"),
     ("Réserves / points en suspens",""),("Nombre d'anomalies bloquantes restantes",""),
     ("Nom du valideur (client)",""),("Fonction",""),("Date de validation",""),
     ("Signature (nom + « bon pour accord »)",""),("Chef de projet Promo.dev","")]:
        l=va.cell(r,2,label); l.font=F(11,True,"26324A"); l.fill=fill(GREY); l.alignment=leftc; bd(va,l.coordinate)
        c=va.cell(r,3); c.fill=fill(YELLOW); c.alignment=Alignment(vertical="center",wrap_text=True,indent=1)
        yb=Side(style="thin",color=YELLOWB); c.border=Border(left=yb,right=yb,top=yb,bottom=yb)
        if val=="__DROPDOWN__":
            c.value="— cliquez pour choisir —"; c.font=F(11,it=True,color="9AA0AA"); valrow=r
        else:
            c.value=val; c.font=F(11)
        va.row_dimensions[r].height=38 if val=="" else 28; r+=1
    dvv=DataValidation(type="list",formula1='"Oui sans réserve,Oui avec réserves,Non"',allow_blank=True)
    dvv.add(f"C{valrow}"); va.add_data_validation(dvv)
    va.conditional_formatting.add(f"C{valrow}",CellIsRule(operator="equal",formula=['"Oui sans réserve"'],fill=fill(GREENF),font=F(11,True,GREENT)))
    va.conditional_formatting.add(f"C{valrow}",CellIsRule(operator="equal",formula=['"Oui avec réserves"'],fill=fill(YELLOW),font=F(11,True,"B45309")))
    va.conditional_formatting.add(f"C{valrow}",CellIsRule(operator="equal",formula=['"Non"'],fill=fill(REDF),font=F(11,True,REDT)))
    va.page_setup.orientation="portrait"; va.page_setup.fitToWidth=1; va.page_setup.fitToHeight=0
    va.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); va.print_area=f"A1:D{r}"


def build_odr(wb):
    # ---------- 1) MODE D'EMPLOI ----------
    ws=wb.active; ws.title="Mode d'emploi"; ws.sheet_view.showGridLines=False
    ws.column_dimensions["A"].width=2.5; ws.column_dimensions["B"].width=40
    ws.column_dimensions["C"].width=76; ws.column_dimensions["D"].width=3
    r=title_bar(ws,2,3,2,"PROCÉDURE DE TEST — OFFRE DE REMBOURSEMENT (ODR)",
                "Gabarit générique · à dupliquer et compléter pour chaque opération"); r+=1
    hc=ws.cell(r,2,"  À LIRE EN PREMIER"); hc.font=F(12,True,WHITE); hc.fill=fill(ACCENT)
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3); hc.alignment=leftc; ws.row_dimensions[r].height=22; r+=1
    intro=("Vous allez tester le site de votre opération (une « offre de remboursement », ou ODR) avant sa mise en ligne. "
           "Ce site « gabarit » est une copie de test, identique à ce que verront les consommateurs. "
           "Objectif : vérifier que tout fonctionne — parcours, formulaire, e-mails — et le valider avant le lancement.\n\n"
           "Comptez environ 30 minutes. Avant de commencer, munissez-vous : d'une boîte mail que vous pouvez consulter "
           "(pensez à vérifier les spams), d'un fichier à téléverser en guise de justificatif (photo ou PDF d'un ticket, "
           "ou toute image de test), et testez de préférence sur ordinateur puis sur mobile. "
           "En cas de blocage, contactez votre chef de projet Promo.dev (coordonnées ci-dessous).")
    ic=ws.cell(r,2,intro); ic.font=F(11); ic.alignment=Alignment(wrap_text=True,vertical="top",indent=1); ic.fill=fill("FBF3E6")
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    for cc in (2,3): ws.cell(r,cc).border=border
    ws.row_dimensions[r].height=150; r+=2
    section(ws,2,r,"1 · Informations de l'opération"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    tag=ws.cell(r,2,"▸ À compléter par le chef de projet Promo.dev avant l'envoi au client")
    tag.font=F(10,True,ACCENT); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    tag.alignment=Alignment(indent=1,vertical="center"); ws.row_dimensions[r].height=20; r+=1
    info=[("Client / Marque","ex. Marque (société)"),("Nom de l'opération","ex. ODR Automne 2026"),
     ("N° opération (idgame)","ex. 999"),("URL de test","ex. https://test.promo.dev/xxxxxxxx"),
     ("IBAN de test à utiliser","ex. FR7630001007941234567890185"),("E-mail de test","ex. prenom.nom+test@domaine.fr"),
     ("Type de justificatif attendu","ex. ticket de caisse / facture (PDF, JPG, PNG)"),("Montant du remboursement","ex. 10 €"),
     ("Dédoublonnage paramétré","ex. 1 participation / e-mail + IBAN"),("Dates de l'offre","ex. 10/10/2026 → 30/11/2026"),
     ("Testeur(s) côté client","ex. Prénom Nom"),("Contact Promo.dev (chef de projet)","ex. prenom@promo.dev"),
     ("Date limite de retour des tests","ex. JJ/MM/AAAA")]
    for label,exv in info:
        lc=ws.cell(r,2,label); lc.font=F(11,True,"26324A"); lc.fill=fill(GREY); lc.alignment=leftc; bd(ws,lc.coordinate)
        val=INFO.get(label,exv); is_real=label in INFO
        vc=ws.cell(r,3,val); vc.font=F(11,False,"26324A" if is_real else "9AA0AA"); vc.fill=fill(YELLOW); vc.alignment=leftc
        yb=Side(style="thin",color=YELLOWB); vc.border=Border(left=yb,right=yb,top=yb,bottom=yb)
        ws.row_dimensions[r].height=26; r+=1
    r+=1
    section(ws,2,r,"2 · Comment utiliser ce classeur"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    steps=["Vérifiez les informations de l'opération ci-dessus (pré-remplies par Promo.dev) : URL, IBAN de test, dates, contact.",
     "Onglet « Guide Houston » : familiarisez-vous avec la barre d'administration et les états de l'offre.",
     "Onglet « Checklist ODR » : déroulez chaque test, indiquez un statut par appareil (OK / KO / N.A.) et un commentaire.",
     "Onglet « Anomalies » : décrivez précisément chaque bug (page, action, capture).",
     "Onglet « Validation » : donnez votre bon pour mise en ligne une fois tous les tests OK.",
     "Renvoyez le fichier complété au chef de projet avant la date limite."]
    for i,s in enumerate(steps,1):
        b=ws.cell(r,2,f"Étape {i}"); b.font=F(11,True,ACCENT); b.alignment=Alignment(vertical="center",indent=1)
        b.fill=fill(ZEBRA if i%2==0 else WHITE)
        c=ws.cell(r,3,s); c.font=F(11); c.alignment=wrapv; c.fill=fill(ZEBRA if i%2==0 else WHITE)
        ws.row_dimensions[r].height=32; r+=1
    r+=1
    section(ws,2,r,"3 · Légende des statuts"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    for code,desc,bg,tx in [("OK","Conforme, rien à signaler",GREENF,GREENT),
     ("KO","Anomalie détectée → à détailler dans l'onglet Anomalies",REDF,REDT),
     ("N.A.","Non applicable à cette opération",NEUTF,"555555"),("À tester","Test pas encore réalisé",YELLOW,"B45309")]:
        p=ws.cell(r,2,code); p.font=F(11,True,tx); p.fill=fill(bg); p.alignment=center; bd(ws,p.coordinate)
        d=ws.cell(r,3,desc); d.font=F(11); d.alignment=leftc; ws.row_dimensions[r].height=22; r+=1
    r+=1
    section(ws,2,r,"4 · Lexique"); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
    ws.cell(r,3).fill=fill(TINT); ws.row_dimensions[r].height=24; r+=1
    lex=[("ODR — Offre de remboursement","Opération où le consommateur est remboursé (en tout ou partie) après un achat, sur présentation d'un justificatif."),
     ("Site gabarit","Version de test du site de l'opération, identique au site final, pour tout valider avant la mise en ligne."),
     ("Houston","Outil interne Promo.dev (la barre en haut du site de test) qui sert à simuler les états de l'offre. Invisible pour les consommateurs."),
     ("Parcours consommateur","Enchaînement des pages suivies par le participant, de l'accueil à la confirmation."),
     ("Justificatif","Preuve d'achat demandée au participant (ticket de caisse, facture…)."),
     ("Dédoublonnage","Contrôle qui empêche les participations en double (même e-mail / IBAN)."),
     ("Quota / compteur","Limite du nombre de participations prévue pour l'offre."),
     ("Bon pour mise en ligne","Validation finale du client avant la mise en ligne.")]
    for term,desc in lex:
        tcell=ws.cell(r,2,term); tcell.font=F(11,True,"26324A"); tcell.fill=fill(GREY)
        tcell.alignment=Alignment(vertical="top",wrap_text=True,indent=1); bd(ws,tcell.coordinate)
        dcell=ws.cell(r,3,desc); dcell.font=F(11); dcell.alignment=Alignment(vertical="top",wrap_text=True,indent=1); bd(ws,dcell.coordinate)
        ws.row_dimensions[r].height=32; r+=1
    ws.page_setup.orientation="portrait"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ws.print_area=f"A1:D{r}"

    # ---------- 2) GUIDE HOUSTON ----------
    g=wb.create_sheet("Guide Houston"); g.sheet_view.showGridLines=False
    g.column_dimensions["A"].width=2.5; g.column_dimensions["B"].width=30
    g.column_dimensions["C"].width=60; g.column_dimensions["D"].width=60
    r=title_bar(g,2,4,2,"GUIDE DE LA BARRE D'ADMINISTRATION « HOUSTON »",
                "La barre en haut du site de test permet de simuler tous les statuts et pages de l'offre."); r+=1
    hh=g.cell(r,2,"⚠  La barre « Houston » est un outil de test interne"); hh.font=F(12,True,WHITE); hh.fill=fill(NAVY)
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); hh.alignment=leftc; g.row_dimensions[r].height=24; r+=1
    frm=("Elle apparaît uniquement sur ce site de test et sert à simuler les différents états de l'offre "
         "(en cours, en attente, terminé…). Elle ne sera JAMAIS visible par les consommateurs sur le site final.")
    fc=g.cell(r,2,frm); fc.font=F(11); fc.alignment=Alignment(wrap_text=True,vertical="center",indent=1); fc.fill=fill("FBF3E6")
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    for cc in range(2,5): g.cell(r,cc).border=border
    g.row_dimensions[r].height=48; r+=2
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    for cc in range(2,5): g.cell(r,cc).fill=fill(TINT); g.cell(r,cc).border=border
    place_abs(g,"image9.png",2,r,720,90,3,90); r+=2
    qh=g.cell(r,2,"  PAR OÙ COMMENCER — testez en 6 étapes"); qh.font=F(12,True,WHITE); qh.fill=fill(ACCENT)
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); qh.alignment=leftc; g.row_dimensions[r].height=22; r+=1
    qs=["Ouvrez l'URL de test (voir l'onglet « Mode d'emploi »).",
        "Dans la barre Houston, réglez « État » sur « En cours ».",
        "Cliquez sur « Je participe » et remplissez le formulaire.",
        "Téléversez un justificatif d'achat et saisissez l'IBAN de test fourni.",
        "Validez : une page de confirmation s'affiche et vous recevez un e-mail.",
        "Notez vos remarques dans « Checklist ODR » (OK/KO) et détaillez tout souci dans « Anomalies »."]
    for i,s in enumerate(qs,1):
        zeb=ZEBRA if i%2==0 else WHITE
        eb=g.cell(r,2,f"Étape {i}"); eb.font=F(11,True,ACCENT); eb.alignment=Alignment(vertical="top",indent=1); eb.fill=fill(zeb); bd(g,eb.coordinate)
        sc=g.cell(r,3,s); sc.font=F(11); sc.alignment=wrap; sc.fill=fill(zeb)
        g.merge_cells(start_row=r,start_column=3,end_row=r,end_column=4); g.cell(r,4).fill=fill(zeb)
        bd(g,f"C{r}"); bd(g,f"D{r}")
        g.row_dimensions[r].height=42; r+=1
    r+=1
    g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    nc=g.cell(r,2,"ℹ Les aperçus ci-dessous sont des exemples génériques. La barre d'administration Houston est identique pour toutes les opérations ; seules les pages de l'offre changent selon la marque.")
    nc.font=F(10,True,"6B7280"); nc.fill=fill("F2F4F8"); nc.alignment=Alignment(wrap_text=True,vertical="center",indent=1)
    for cc in range(2,5): g.cell(r,cc).border=border
    g.row_dimensions[r].height=32; r+=2
    hr=r
    for i,txt in enumerate(("CONTRÔLE","À QUOI ÇA SERT","APERÇU")):
        c=g.cell(hr,2+i,txt); c.font=F(11,True,WHITE); c.fill=fill(BLUE); c.alignment=wrapc; bd(g,c.coordinate)
    g.row_dimensions[hr].height=24
    guide=[("CONNEXION","Se connecter sur l'URL de test de l'opération (voir onglet « Mode d'emploi »). Un IBAN de test est fourni pour remplir le formulaire.",None,0),
     ("ÉTAT","Permet de tester les différents statuts de l'opération (menu déroulant).","image6.png",300),
     ("ÉTAT ▸ En cours","Le parcours consommateur est accessible ; le bouton « Je participe » apparaît pour lancer et valider une participation.","image1.png",210),
     ("ÉTAT ▸ En attente","Affiche la page d'attente : l'offre n'a pas encore commencé.","neutral_attente.png",380),
     ("ÉTAT ▸ Terminé","Affiche la page de fin : l'offre est terminée.","neutral_termine.png",380),
     ("ÉTAT ▸ Offline","Affiche la page « hors ligne » quand l'offre est suspendue.","image12.png",380),
     ("ÉTAT ▸ Quota atteint","Affiche la page d'accueil quand le nombre max de participations est atteint (si compteur limitatif).","image13.png",380),
     ("VISUEL","Permet de visualiser au choix la Home Page ou la page de confirmation de participation.","image7.png",400),
     ("PARTICIPATION (statut)","PERD ou GAGNE. Sur une ODR, sans incidence sur le parcours ni la page de confirmation (ce n'est pas un jeu mais un remboursement).","image8.png",240),
     ("DÉSACTIVER CONTRÔLES","Désactive le contrôle d'unicité : permet de retester plusieurs fois avec le même e-mail / IBAN. Décoché = le dédoublonnage paramétré s'applique.","image3.png",260),
     ("SUIVI DE PARTICIPATION","Depuis le pied de page, permet de recevoir un e-mail avec le lien de suivi (statut de traitement du dossier).","neutral_suivi.png",360)]
    row=hr+1
    for idx,(label,desc,imgn,tw) in enumerate(guide):
        zeb=ZEBRA if idx%2==0 else WHITE
        lc=g.cell(row,2,label); lc.font=F(11,True,NAVY); lc.alignment=wrapv; lc.fill=fill(TINT); bd(g,lc.coordinate)
        dc=g.cell(row,3,desc); dc.font=F(12); dc.alignment=wrap; dc.fill=fill(zeb); bd(g,dc.coordinate)
        ac=g.cell(row,4); ac.fill=fill(zeb); bd(g,ac.coordinate)
        if imgn: place_abs(g,imgn,4,row,tw,230,1,60)
        else: g.row_dimensions[row].height=42
        row+=1
    g.page_setup.orientation="landscape"; g.page_setup.fitToWidth=1; g.page_setup.fitToHeight=0
    g.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); g.print_area=f"A1:D{row}"

    # ---------- 3) CHECKLIST ODR ----------
    ck=wb.create_sheet("Checklist ODR"); ck.sheet_view.showGridLines=False
    for col,w in {"A":2.5,"B":5.5,"C":23,"D":46,"E":46,"F":12,"G":12,"H":12,"I":34}.items(): ck.column_dimensions[col].width=w
    r=title_bar(ck,2,9,2,"CHECKLIST DE TEST — ODR",
                "Renseignez un statut par appareil (Desktop / Tablette / Mobile). Tout KO doit être détaillé dans l'onglet « Anomalies »."); r+=1
    headers=["N°","Catégorie","Test à réaliser","Résultat attendu","Desktop","Tablette","Mobile","Commentaire / Anomalie"]
    hr=r
    for i,h in enumerate(headers):
        c=ck.cell(hr,2+i,h); c.font=F(11,True,WHITE); c.fill=fill(BLUE); c.alignment=wrapc; bd(ck,c.coordinate)
    ck.row_dimensions[hr].height=32
    LC={"Accès & états":"E7EEF9","Home page":"E9F5EC","Formulaire":"FBF0E4","Justificatif":"F3EAF7",
     "IBAN / RIB":"E4F3F5","Consentements":"FDECEC","Validation":"EAF0FB","E-mails":"FFF7E0","Anti-fraude":"F0ECF9"}
    tests=[("Accès & états","Accéder à l'URL de test","La page se charge, la barre Houston est visible en haut"),
     ("Accès & états","État « En attente »","La page d'attente s'affiche (offre pas encore commencée, date annoncée)"),
     ("Accès & états","État « En cours »","Le bouton « Je participe » apparaît, le parcours est accessible"),
     ("Accès & états","État « Terminé »","La page de fin d'offre s'affiche"),
     ("Accès & états","État « Offline »","La page « opération suspendue » s'affiche"),
     ("Accès & états","État « Quota atteint » (si compteur)","La page « limite de participations atteinte » s'affiche"),
     ("Home page","Affichage Home desktop","Visuels, logos, textes et charte conformes à la maquette validée"),
     ("Home page","Affichage Home mobile (responsive)","Mise en page correcte sur smartphone, aucun élément coupé"),
     ("Home page","Mentions légales / règlement / confidentialité","Liens présents et documents accessibles"),
     ("Home page","Bouton « Je participe »","Lance le parcours de participation"),
     ("Formulaire","Présence de tous les champs attendus","Tous les champs du cahier des charges sont présents"),
     ("Formulaire","Champs obligatoires laissés vides","Message d'erreur clair, blocage de la validation"),
     ("Formulaire","Contrôles de format (e-mail, code postal, tél.)","Une saisie invalide déclenche un message d'erreur"),
     ("Justificatif","Upload d'un justificatif valide (PDF/JPG/PNG)","Le fichier est accepté et visible dans le récap"),
     ("Justificatif","Upload d'un format/poids non autorisé","Message d'erreur, fichier refusé"),
     ("IBAN / RIB","Saisie de l'IBAN de test (valide)","L'IBAN est accepté"),
     ("IBAN / RIB","Saisie d'un IBAN invalide","Message d'erreur, blocage de la validation"),
     ("Consentements","Case CGU / RGPD non cochée","Impossible de valider tant que le consentement obligatoire n'est pas donné"),
     ("Validation","Récapitulatif avant envoi","Les données saisies sont exactes et modifiables"),
     ("Validation","Valider la participation","Page de confirmation affichée (visuel + message de succès)"),
     ("E-mails","E-mail de confirmation de participation","E-mail reçu ; expéditeur, objet, visuels et liens corrects"),
     ("E-mails","Suivi de participation (pied de page)","Après saisie de l'e-mail, réception du lien de suivi"),
     ("E-mails","Page de suivi de participation","Le bon statut de traitement du dossier s'affiche"),
     ("Anti-fraude","Dédoublonnage (contrôles activés)","2e participation même e-mail/IBAN refusée avec message"),
     ("Anti-fraude","Désactiver contrôles","Permet de re-tester avec le même e-mail / IBAN")]
    row=hr+1
    ex=("Ex","Formulaire","Laisser le champ « e-mail » vide et valider","Message « Champ obligatoire », validation bloquée","OK","OK","KO","KO sur mobile : bouton coupé")
    for i,v in enumerate(ex):
        c=ck.cell(row,2+i,v); c.border=border; c.alignment=(wrapc if i in(0,4,5,6) else wrap); c.font=F(10,it=True,color="9AA0AA"); c.fill=fill("FBFBFD")
    ck.row_dimensions[row].height=28; row+=1
    start=row; n=1
    for cat,test,exp in tests:
        zeb=ZEBRA if n%2==0 else WHITE
        a=ck.cell(row,2,n); a.alignment=center; a.font=F(11,True,"6B7280"); a.fill=fill(zeb); bd(ck,a.coordinate)
        cc=ck.cell(row,3,cat); cc.font=F(11,True,"2C3E63"); cc.alignment=wrapv; cc.fill=fill(LC.get(cat,"EDEFF3")); bd(ck,cc.coordinate)
        t=ck.cell(row,4,test); t.font=F(11); t.alignment=wrap; t.fill=fill(zeb); bd(ck,t.coordinate)
        e=ck.cell(row,5,exp); e.font=F(11,color="4B5563"); e.alignment=wrap; e.fill=fill(zeb); bd(ck,e.coordinate)
        for col in (6,7,8):
            x=ck.cell(row,col); x.fill=fill(zeb); x.alignment=center; bd(ck,x.coordinate); x.font=F(11)
        cm=ck.cell(row,9); cm.fill=fill(zeb); cm.alignment=wrap; bd(ck,cm.coordinate); cm.font=F(11)
        ck.row_dimensions[row].height=34; row+=1; n+=1
    end=row-1
    dv=DataValidation(type="list",formula1='"OK,KO,N.A.,À tester"',allow_blank=True); dv.add(f"F{start}:H{end}"); ck.add_data_validation(dv)
    for val,bg,tx in [("OK",GREENF,GREENT),("KO",REDF,REDT),("N.A.",NEUTF,"555555"),("À tester",YELLOW,"B45309")]:
        ck.conditional_formatting.add(f"F{start}:H{end}",CellIsRule(operator="equal",formula=[f'"{val}"'],fill=fill(bg),font=F(11,True,tx)))
    sr=end+2
    ck.merge_cells(start_row=sr,start_column=3,end_row=sr,end_column=9)
    hh=ck.cell(sr,3,"  RÉCAPITULATIF (nombre de tests par statut et par appareil)"); hh.font=F(12,True,WHITE); hh.fill=fill(NAVY); hh.alignment=leftc; ck.row_dimensions[sr].height=24
    ck.merge_cells(start_row=sr+1,start_column=3,end_row=sr+1,end_column=5)
    for cidx,dev in [(6,"Desktop"),(7,"Tablette"),(8,"Mobile")]:
        dc=ck.cell(sr+1,cidx,dev); dc.font=F(11,True,NAVY); dc.fill=fill(TINT); dc.alignment=center; bd(ck,dc.coordinate)
    rr=sr+2
    for lab,kind,bg,tx in [("Tests OK","OK",GREENF,GREENT),("Tests KO","KO",REDF,REDT),("Restant à tester","REST",YELLOW,"B45309")]:
        ck.merge_cells(start_row=rr,start_column=3,end_row=rr,end_column=5)
        l=ck.cell(rr,3,lab); l.font=F(11,True,tx); l.fill=fill(bg); l.alignment=leftc; bd(ck,l.coordinate)
        for cidx,cl in [(6,"F"),(7,"G"),(8,"H")]:
            if kind=="REST":
                form=f'=COUNTA(D{start}:D{end})-COUNTIF({cl}{start}:{cl}{end},"OK")-COUNTIF({cl}{start}:{cl}{end},"KO")-COUNTIF({cl}{start}:{cl}{end},"N.A.")'
            else:
                form=f'=COUNTIF({cl}{start}:{cl}{end},"{kind}")'
            v=ck.cell(rr,cidx,form); v.font=F(12,True,tx); v.fill=fill(bg); v.alignment=center; bd(ck,v.coordinate)
        ck.row_dimensions[rr].height=22; rr+=1
    ck.freeze_panes=f"B{hr+1}"
    ck.page_setup.orientation="landscape"; ck.page_setup.fitToWidth=1; ck.page_setup.fitToHeight=0
    ck.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ck.print_area=f"A1:I{rr}"
    build_anomalies(wb); build_validation(wb)


# ==================== DISPATCH ====================
# build(game,oblig) [jeux] crée et renvoie son propre wb ; build_odr/prime/formulaire
# construisent dans un wb passé. On uniformise en callables 0-arg -> wb.
def _fresh(builder):
    wb=openpyxl.Workbook(); builder(wb); return wb
MECHANICS={
 "ODR":        lambda: _fresh(build_odr),
 "IG_AOA":     lambda: build("IG","AOA"),
 "IG_SOA":     lambda: build("IG","SOA"),
 "TAS_AOA":    lambda: build("TAS","AOA"),
 "TAS_SOA":    lambda: build("TAS","SOA"),
 "AUTO_AOA":   lambda: build("AUTO","AOA"),
 "AUTO_SOA":   lambda: build("AUTO","SOA"),
 "PRIME":      lambda: _fresh(build_prime),
 "FORMULAIRE": lambda: _fresh(build_formulaire),
}

def main():
    import argparse, json
    ap=argparse.ArgumentParser(description="Génère la procédure de test Promo.dev adaptée à la mécanique.")
    ap.add_argument("--mechanic",required=True,choices=list(MECHANICS.keys()))
    ap.add_argument("--info",help="Chemin d'un JSON { 'libellé exact onglet 1': 'valeur' } (facultatif)")
    ap.add_argument("--out",required=True,help="Chemin du .xlsx de sortie")
    a=ap.parse_args()
    global INFO
    if a.info:
        with open(a.info,encoding="utf-8") as f: INFO=json.load(f)
    wb=MECHANICS[a.mechanic]()
    wb.save(a.out)
    print("OK:",a.out)

if __name__=="__main__":
    main()
